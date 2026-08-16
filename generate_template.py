#!/usr/bin/env python3
"""Generate fiber circuit + port-link field entry template."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from fiber_infra import port_name_for_cable_fiber
from site_devices import DEV_BHY1, DEV_BHY2, DEV_SJ1, DEV_SJ2, DEV_ZG

OUTPUT = Path(__file__).with_name("现场登记表.xlsx")
OUTPUT_LEGACY = Path(__file__).with_name("现场光路录入.xlsx")

# 得丰焦化默认坐标（与 NetBox OSP 地图中心一致）
DEFAULT_SITE = "得丰焦化"
DEFAULT_LAT = "41.793545"
DEFAULT_LON = "119.264694"

DATA_HEADERS = [
    "光路编号*",
    "段序号",
    "光路名称",
    "缆段编号",
    "A端设备*",
    "A端纤芯(排-芯)*",
    "A端端口名(参考)",
    "B端设备*",
    "B端纤芯(排-芯)*",
    "B端端口名(参考)",
    "业务说明",
    "IP地址",
    "实测光损(dB)",
    "测试波长(nm)",
    "标签位置",
    "扫码链接",
    "二维码",
]
DATA_WIDTHS = [16, 8, 18, 20, 34, 14, 22, 34, 14, 22, 22, 14, 12, 12, 36, 46, 14]

PORT_LINK_HEADERS = [
    "A端设备*",
    "A端缆段编号",
    "A端纤芯(排-芯)*",
    "A端端口名(参考)",
    "B端设备*",
    "B端缆段编号",
    "B端纤芯(排-芯)*",
    "B端端口名(参考)",
    "备注",
    "关联光路编号",
]
PORT_LINK_WIDTHS = [34, 20, 14, 22, 34, 20, 14, 22, 28, 16]

def port_link_row(
    dev_a: str,
    cab_a: str,
    port_a: str,
    dev_b: str,
    cab_b: str,
    port_b: str,
    note: str = "",
    cid: str = "",
) -> list:
    pa = port_name_for_cable_fiber(cab_a, port_a)
    pb = port_name_for_cable_fiber(cab_b, port_b)
    if not note:
        note = f"{pa} ↔ {pb}"
    return [dev_a, cab_a, port_a, pa, dev_b, cab_b, port_b, pb, note, cid]


PORT_LINK_EXAMPLE_ROWS = [
    port_link_row(
        DEV_BHY1, "CBL-轧钢-白灰窑", "1-1", DEV_BHY1, "CBL-白灰窑-白灰窑2", "1-1",
        note="监控：轧钢至白灰窑1-1 ↔ 白灰窑至白灰窑21-1", cid="FIB-轧钢-监控",
    ),
    port_link_row(
        DEV_BHY1, "CBL-轧钢-白灰窑", "2-2", DEV_BHY1, "CBL-白灰窑-烧结", "2-2",
        note="主控：轧钢至白灰窑2-2 ↔ 白灰窑至烧结2-2", cid="FIB-轧钢-主控",
    ),
    port_link_row(
        DEV_BHY2, "CBL-白灰窑-白灰窑2", "1-1", DEV_BHY2, "CBL-白灰窑2-烧结", "1-1",
        note="监控：白灰窑至白灰窑21-1 ↔ 白灰窑2至烧结1-1", cid="FIB-轧钢-监控",
    ),
]

SPLICE_HEADERS = [
    "熔接点设备*",
    "熔接方案名称",
    "A端缆段编号",
    "A端纤芯(排-芯)*",
    "A端端口名(参考)",
    "B端缆段编号",
    "B端纤芯(排-芯)*",
    "B端端口名(参考)",
    "托盘序号",
    "备注",
    "关联光路编号",
]
SPLICE_WIDTHS = [34, 24, 20, 14, 22, 20, 14, 22, 10, 28, 16]

STRAND_LOSS_HEADERS = [
    "缆段编号*",
    "纤芯(排-芯)*",
    "实测光损(dB)*",
    "测试波长(nm)",
    "备注",
    "关联光路编号",
]
STRAND_LOSS_WIDTHS = [22, 14, 14, 12, 24, 16]

STRAND_LOSS_EXAMPLE_ROWS = [
    ["CBL-轧钢-白灰窑", "1-1", "0.35", "1310", "OTDR A→B", "FIB-轧钢-监控"],
    ["CBL-轧钢-白灰窑", "2-2", "0.38", "1310", "OTDR A→B", "FIB-轧钢-主控"],
    ["CBL-白灰窑-白灰窑2", "1-1", "0.42", "1310", "", "FIB-轧钢-监控"],
    ["CBL-白灰窑2-烧结", "1-1", "0.55", "1310", "", "FIB-轧钢-监控"],
    ["CBL-白灰窑-烧结", "2-2", "0.48", "1310", "", "FIB-轧钢-主控"],
    ["CBL-烧结-白灰窑", "1-1", "0.40", "1550", "", "FIB-烧结-SCADA"],
    ["CBL-白灰窑-轧钢", "1-1", "0.36", "1550", "", "FIB-烧结-SCADA"],
]

STRAND_LOSS_HEADER_FILL = PatternFill("solid", fgColor="C55A11")

SITE_HEADERS = [
    "站点名称*",
    "URL标识(slug)*",
    "状态",
    "物理地址",
    "纬度",
    "经度",
    "联系人",
    "备注",
]
SITE_WIDTHS = [18, 18, 10, 28, 14, 14, 12, 24]
SITE_HEADER_FILL = PatternFill("solid", fgColor="375623")

DEVICE_HEADERS = [
    "站点名称*",
    "机房/位置*",
    "设备名称*",
    "设备角色",
    "设备型号",
    "机柜号",
    "U位",
    "纬度",
    "经度",
    "状态",
    "备注",
]
DEVICE_WIDTHS = [16, 20, 34, 10, 12, 10, 8, 14, 14, 10, 24]
DEVICE_HEADER_FILL = PatternFill("solid", fgColor="548235")

TRUNK_HEADERS = [
    "缆段编号*",
    "A端设备*",
    "B端设备*",
    "缆段长度(m)",
    "长度单位",
    "光纤型号*",
    "芯数*",
    "光缆结构",
    "是否铠装",
    "A端后端口名(系统)",
    "B端后端口名(系统)",
    "备注",
]
TRUNK_WIDTHS = [22, 34, 34, 12, 10, 14, 8, 12, 10, 24, 24, 24]
TRUNK_HEADER_FILL = PatternFill("solid", fgColor="BF8F00")

SYSTEM_MAP_HEADERS = [
    "NetBox 对象",
    "系统字段(API名)",
    "Excel 工作表",
    "Excel 列名",
    "必填",
    "创建方式",
    "说明",
]
SYSTEM_MAP_WIDTHS = [22, 22, 16, 22, 8, 12, 40]

SITE_EXAMPLE_ROWS = [
    [DEFAULT_SITE, "defeng-coking", "active", "赤峰宁城汐子工业园", DEFAULT_LAT, DEFAULT_LON, "", "厂区主站点"],
]

DEVICE_EXAMPLE_ROWS = [
    [DEFAULT_SITE, "信通机房", DEV_ZG, "ODF", "ODF-48", "JG01", "1", DEFAULT_LAT, DEFAULT_LON, "active", "轧钢上行ODF"],
    [DEFAULT_SITE, "白灰窑机房", DEV_BHY1, "ODF", "ODF-48", "A01", "1", "", "", "active", ""],
    [DEFAULT_SITE, "白灰窑机房", DEV_BHY2, "ODF", "ODF-48", "A01", "2", "", "", "active", ""],
    [DEFAULT_SITE, "烧结主控机房", DEV_SJ1, "ODF", "ODF-48", "A01", "1", "", "", "active", ""],
    [DEFAULT_SITE, "烧结主控机房", DEV_SJ2, "ODF", "ODF-48", "A01", "2", "", "", "active", ""],
]

TRUNK_EXAMPLE_ROWS = [
    ["CBL-轧钢-白灰窑", DEV_ZG, DEV_BHY1, "850", "m", "G.652D-12", "12", "松套层绞", "否", "RP-CBL-轧钢-白灰窑-A", "RP-CBL-轧钢-白灰窑-B", ""],
    ["CBL-白灰窑-白灰窑2", DEV_BHY1, DEV_BHY2, "120", "m", "G.652D-12", "12", "松套层绞", "否", "", "", "站内第二框"],
    ["CBL-白灰窑2-烧结", DEV_BHY2, DEV_SJ1, "1200", "m", "G.652D-12", "12", "松套层绞", "否", "", "", ""],
    ["CBL-白灰窑-烧结", DEV_BHY1, DEV_SJ1, "1100", "m", "G.652D-12", "12", "松套层绞", "否", "", "", ""],
    ["CBL-烧结-白灰窑", DEV_SJ2, DEV_BHY1, "1150", "m", "G.652D-12", "12", "松套层绞", "否", "", "", ""],
    ["CBL-白灰窑-轧钢", DEV_BHY1, DEV_ZG, "900", "m", "G.652D-12", "12", "松套层绞", "否", "", "", ""],
]


def circuit_row(
    cid: str,
    cable: str,
    dev_a: str,
    port_a: str,
    dev_b: str,
    port_b: str,
    *,
    seg: str = "",
    name: str = "",
    service: str = "",
    ip: str = "",
    loss_db: str = "",
    wavelength: str = "",
    label_pos: str = "",
) -> list:
    pa = port_name_for_cable_fiber(cable, port_a)
    pb = port_name_for_cable_fiber(cable, port_b)
    if not label_pos:
        label_pos = f"{dev_a} {pa}"
    return [
        cid, seg, name, cable, dev_a, port_a, pa, dev_b, port_b, pb,
        service, ip, loss_db, wavelength, label_pos, "", "",
    ]


def splice_row(
    dev: str,
    plan: str,
    cab_a: str,
    port_a: str,
    cab_b: str,
    port_b: str,
    *,
    tray: str = "1",
    note: str = "",
    cid: str = "",
) -> list:
    pa = port_name_for_cable_fiber(cab_a, port_a)
    pb = port_name_for_cable_fiber(cab_b, port_b)
    if not note:
        note = f"{pa} ↔ {pb}"
    return [dev, plan, cab_a, port_a, pa, cab_b, port_b, pb, tray, note, cid]


EXAMPLE_ROWS = [
    circuit_row(
        "FIB-轧钢-监控", "CBL-轧钢-白灰窑",
        DEV_ZG, "1-1", DEV_BHY1, "1-1",
        name="轧钢线监控", service="轧钢→白灰窑", ip="10.9.30.31",
    ),
    circuit_row(
        "FIB-轧钢-监控", "CBL-白灰窑-白灰窑2",
        DEV_BHY1, "1-1", DEV_BHY2, "1-1",
        service="白灰窑站内跳接",
    ),
    circuit_row(
        "FIB-轧钢-监控", "CBL-白灰窑2-烧结",
        DEV_BHY2, "1-1", DEV_SJ1, "1-1",
        service="白灰窑→烧结", loss_db="2.80", wavelength="1310",
        label_pos=f"{DEV_SJ1} 白灰窑2至烧结1-1",
    ),
    circuit_row(
        "FIB-轧钢-主控", "CBL-轧钢-白灰窑",
        DEV_ZG, "2-2", DEV_BHY1, "2-2",
        name="轧钢主控上传", service="轧钢→白灰窑", ip="10.9.30.32",
    ),
    circuit_row(
        "FIB-轧钢-主控", "CBL-白灰窑-烧结",
        DEV_BHY1, "2-2", DEV_SJ1, "2-2",
        service="白灰窑→烧结", loss_db="2.75", wavelength="1310",
        label_pos=f"{DEV_SJ1} 白灰窑至烧结2-2",
    ),
    circuit_row(
        "FIB-烧结-SCADA", "CBL-烧结-白灰窑",
        DEV_SJ2, "1-1", DEV_BHY1, "1-1",
        name="烧结SCADA采集", service="烧结→白灰窑", ip="10.9.20.12",
    ),
    circuit_row(
        "FIB-烧结-SCADA", "CBL-白灰窑-轧钢",
        DEV_BHY1, "1-1", DEV_ZG, "1-1",
        service="白灰窑→轧钢", loss_db="2.68", wavelength="1550",
        label_pos=f"{DEV_ZG} 白灰窑至轧钢1-1",
    ),
]

SPLICE_EXAMPLE_ROWS = [
    splice_row(
        DEV_BHY1, f"{DEV_BHY1}-熔接方案",
        "CBL-轧钢-白灰窑", "1-1", "CBL-白灰窑-白灰窑2", "1-1",
        note="监控：轧钢至白灰窑1-1 ↔ 白灰窑至白灰窑21-1", cid="FIB-轧钢-监控",
    ),
    splice_row(
        DEV_BHY1, f"{DEV_BHY1}-熔接方案",
        "CBL-轧钢-白灰窑", "2-2", "CBL-白灰窑-烧结", "2-2",
        note="主控：轧钢至白灰窑2-2 ↔ 白灰窑至烧结2-2", cid="FIB-轧钢-主控",
    ),
    splice_row(
        DEV_BHY2, f"{DEV_BHY2}-熔接方案",
        "CBL-白灰窑-白灰窑2", "1-1", "CBL-白灰窑2-烧结", "1-1",
        note="监控：白灰窑至白灰窑21-1 ↔ 白灰窑2至烧结1-1", cid="FIB-轧钢-监控",
    ),
    splice_row(
        DEV_BHY2, f"{DEV_BHY2}-熔接方案",
        "CBL-白灰窑-白灰窑2", "2-2", "CBL-白灰窑2-烧结", "2-2",
        note="备用：白灰窑至白灰窑22-2 ↔ 白灰窑2至烧结2-2",
    ),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
SPLICE_HEADER_FILL = PatternFill("solid", fgColor="7030A0")
HEADER_FONT = Font(color="FFFFFF", bold=True)
NOTE_FONT = Font(color="666666", italic=True)
SECTION_FILL = PatternFill("solid", fgColor="D9E2F3")
THIN = Side(style="thin", color="CCCCCC")
LAST_COL = get_column_letter(len(DATA_HEADERS))
PORT_LINK_LAST_COL = get_column_letter(len(PORT_LINK_HEADERS))
SPLICE_LAST_COL = get_column_letter(len(SPLICE_HEADERS))
STRAND_LOSS_LAST_COL = get_column_letter(len(STRAND_LOSS_HEADERS))
STRAND_LOSS_HEADER_FILL = PatternFill("solid", fgColor="C55A11")


def style_table_header(ws, row: int, headers: list[str], widths: list[int], fill: PatternFill = HEADER_FILL) -> None:
    for col, (title, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.fill = fill
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width


def write_field_guide(ws) -> None:
    ws.title = "字段说明"

    ws["A1"] = "现场登记表 — 字段定义与导入说明"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:F1")

    ws["A2"] = (
        "工作表顺序：字段说明 → 系统字段对照 → 厂区站点/ODF/缆段（信通预建）→ 光路录入 → 端口连接 → 缆段纤芯光损。"
        "填写后运行 import_from_excel.py --link-mode port 导入业务链路。"
    )
    ws["A2"].font = NOTE_FONT
    ws.merge_cells("A2:F2")

    ws["A4"] = "〇、导入流程（重要）"
    ws["A4"].font = Font(bold=True, size=12)
    ws["A4"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws.merge_cells("A4:F4")
    steps = [
        "1. 信通预建：「厂区站点登记」「ODF设备登记」「厂区缆段登记」→ NetBox DCIM+FMS（可运行 bootstrap_test_infra.py）",
        "2. 现场填写「光路录入」：每段一行，相同光路编号=同一条逻辑链路",
        "3. 填写「端口连接」：同机房 ODF 前置端口直连（FrontPort↔FrontPort），可自动推断或手工补充",
        "4. 填写「缆段纤芯光损」：每条 CBL 缆段、每根纤芯的 OTDR 实测值",
        "5. 运行 python import_from_excel.py → 自动配置端口、跳线、光路路径，并校验 is_complete",
        "6. 可选：python sync_osp_strands.py 同步 OSP 地图追踪数据",
    ]
    for i, text in enumerate(steps, start=5):
        ws[f"A{i}"] = text
        ws.merge_cells(f"A{i}:F{i}")

    ws["A10"] = "一、光路录入 — 多段逻辑链路"
    ws["A10"].font = Font(bold=True, size=12)
    ws["A10"].fill = SECTION_FILL
    ws.merge_cells("A10:F10")
    ws["A11"] = (
        "相同「光路编号」填几行=几段。段序号可留空，导入时按 A/B 衔接自动排序。"
        "「缆段编号」填 NetBox 中 DCIM 缆线 label，留空则导入时自动匹配 A/B 设备间的缆段。"
    )
    ws["A11"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A11:F11")
    ws.row_dimensions[11].height = 36

    ws["A13"] = "二、端口连接 — 站内/跨框 FrontPort 跳接"
    ws["A13"].font = Font(bold=True, size=12)
    ws["A13"].fill = SECTION_FILL
    ws.merge_cells("A13:F13")
    ws["A14"] = (
        "可选。导入默认用 NetBox DCIM「FrontPort ↔ FrontPort」连接替代 FMS 熔接条目。"
        "光路表相邻两段在同一 ODF 换缆、或跨两框延续同一纤芯时，程序会自动补连接，一般不必手工填写。"
        "同机房两 ODF 用跳纤直连时，可在此表填写跨框连接（也可保留光路中间缆段）。"
    )
    ws["A14"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A14:F14")
    ws.row_dimensions[14].height = 48

    ws["A16"] = "三、命名规范"
    ws["A16"].font = Font(bold=True, size=12)
    ws["A16"].fill = SECTION_FILL
    ws.merge_cells("A16:F16")
    naming = [
        ("对象", "格式", "示例"),
        ("ODF设备", "{机房}-{机柜}-ODF-{序号}", "白灰窑机房-A01-ODF-01"),
        ("纤芯（Excel填）", "{排}-{芯}", "2-2"),
        ("缆段编号", "CBL-{起点}-{终点}", "CBL-白灰窑-烧结"),
        ("ODF端口（导入后）", "{起点}至{终点}{排}-{芯}", "白灰窑至烧结2-2"),
        ("光路编号", "FIB-{业务}", "FIB-轧钢-监控"),
    ]
    site_odfs = [
        ("站点", "ODF 设备全名（Excel A/B端设备）"),
        ("轧钢", DEV_ZG),
        ("白灰窑 #1", DEV_BHY1),
        ("白灰窑 #2", DEV_BHY2),
        ("烧结 #1", DEV_SJ1),
        ("烧结 #2", DEV_SJ2),
    ]
    for r, row in enumerate(naming, start=17):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            if r == 17:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT

    ws["A23"] = "三（续）、示例 ODF 设备清单"
    ws["A23"].font = Font(bold=True, size=11)
    ws.merge_cells("A23:C23")
    for r, row in enumerate(site_odfs, start=24):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            if r == 24:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT

    ws["A31"] = "四、光路录入 — 全部字段"
    ws["A31"].font = Font(bold=True, size=12)
    ws["A31"].fill = SECTION_FILL
    ws.merge_cells("A31:F31")
    ws["A32"] = (
        "带 * 为必填。A/B端端口名(参考) 列仅演示导入后 NetBox 端口命名，导入程序不读取；"
        "现场只需填 排-芯 + 缆段编号。A/B端设备必须与 NetBox 设备名完全一致（见上表）。"
    )
    ws["A32"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A32:F32")
    ws.row_dimensions[32].height = 36
    style_table_header(ws, 33, ["字段", "必填", "说明", "示例"], [18, 8, 44, 28])
    circuit_fields = [
        ("光路编号", "是", "相同编号=同一条逻辑链路的多段", "FIB-轧钢-监控"),
        ("段序号", "否", "可留空；填了则按序号排", "留空"),
        ("光路名称", "首段填", "仅该光路第一段填写", "轧钢线监控"),
        ("缆段编号", "建议填", "NetBox DCIM 缆线 label；留空自动匹配 A/B 设备间缆段", "CBL-白灰窑-烧结"),
        ("A端设备", "是", "本段起点 ODF，全名与 NetBox 一致", DEV_ZG),
        ("A端纤芯(排-芯)", "是", "本段起点纤芯，格式 排-芯", "2-2"),
        ("A端端口名(参考)", "否", "导入后自动生成，勿手工改", "轧钢至白灰窑2-2"),
        ("B端设备", "是", "本段终点 ODF", DEV_BHY1),
        ("B端纤芯(排-芯)", "是", "本段终点纤芯；同段通常与 A 端相同", "2-2"),
        ("B端端口名(参考)", "否", "导入后自动生成", "轧钢至白灰窑2-2"),
        ("业务说明", "否", "本段路由说明", "白灰窑→烧结"),
        ("IP地址", "否", "业务 IP，通常填首段", "10.9.30.32"),
        ("实测光损(dB)", "末段填", "全链路端到端 OTDR；与「缆段纤芯光损」表分工", "2.75"),
        ("测试波长(nm)", "否", "配合全链路光损，默认 1310", "1310"),
        ("标签位置", "否", "现场物理标签位置描述", "烧结主控… 白灰窑至烧结2-2"),
        ("扫码链接", "否", "导入后自动回填 NetBox 追踪页 URL", "留空"),
        ("二维码", "否", "导入后自动嵌入 PNG 图片（扫开追踪页）", "留空"),
    ]
    for i, row in enumerate(circuit_fields, start=34):
        for c, val in enumerate(row, start=1):
            ws.cell(row=i, column=c, value=val).border = Border(
                left=THIN, right=THIN, top=THIN, bottom=THIN
            )

    ws["A52"] = "五、端口连接 — 全部字段"
    ws["A52"].font = Font(bold=True, size=12)
    ws["A52"].fill = SECTION_FILL
    ws.merge_cells("A52:F52")
    ws["A53"] = (
        "可选表。光路相邻段换缆/跨框时导入可自动推断；"
        "端口名(参考) 列同样仅演示，填 排-芯 + 缆段编号即可。"
        "A/B 端设备相同=同机房内跳接；不同=跨框 FrontPort 连接。"
    )
    ws["A53"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A53:F53")
    ws.row_dimensions[53].height = 40
    style_table_header(ws, 54, ["字段", "必填", "说明", "示例"], [18, 8, 44, 28])
    port_link_fields = [
        ("A端设备", "是", "连接起点 ODF，全名与 NetBox 一致", DEV_BHY1),
        ("A端缆段编号", "建议填", "A 侧端口命名依据的缆段 label", "CBL-轧钢-白灰窑"),
        ("A端纤芯(排-芯)", "是", "A 侧 排-芯", "1-1"),
        ("A端端口名(参考)", "否", "导入后 A 侧端口名", "轧钢至白灰窑1-1"),
        ("B端设备", "是", "连接终点 ODF", DEV_BHY1),
        ("B端缆段编号", "建议填", "B 侧端口命名依据的缆段 label", "CBL-白灰窑-白灰窑2"),
        ("B端纤芯(排-芯)", "是", "B 侧 排-芯", "1-1"),
        ("B端端口名(参考)", "否", "导入后 B 侧端口名", "白灰窑至白灰窑21-1"),
        ("备注", "否", "连接说明", "轧钢至白灰窑1-1 ↔ 白灰窑至白灰窑21-1"),
        ("关联光路编号", "否", "便于对账，不影响导入", "FIB-轧钢-监控"),
    ]
    for i, row in enumerate(port_link_fields, start=55):
        for c, val in enumerate(row, start=1):
            ws.cell(row=i, column=c, value=val).border = Border(
                left=THIN, right=THIN, top=THIN, bottom=THIN
            )

    ws["A67"] = "六、缆段纤芯光损 — 每缆每芯 OTDR 实测"
    ws["A67"].font = Font(bold=True, size=12)
    ws["A67"].fill = PatternFill("solid", fgColor="FCE4D6")
    ws.merge_cells("A67:F67")
    ws["A68"] = (
        "一行 = 一条物理缆段（CBL-xxx）上的一根纤芯。导入后写入 NetBox："
        "纤缆管理 → 光缆 → 对应缆段 → 纤芯列表 →「实测光损(dB)」。"
        "「关联光路编号」仅便于对账，不影响导入。"
    )
    ws["A68"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A68:F68")
    ws.row_dimensions[68].height = 40
    style_table_header(ws, 69, ["字段", "必填", "说明", "示例"], [16, 8, 44, 24], STRAND_LOSS_HEADER_FILL)
    strand_fields = [
        ("缆段编号", "是", "NetBox DCIM 缆线 label，与光路录入一致", "CBL-烧结-白灰窑"),
        ("纤芯(排-芯)", "是", "排-芯；导入后 ODF 端口如 白灰窑至烧结2-2", "2-2"),
        ("实测光损(dB)", "是", "该缆段该芯 OTDR 单向或往返损耗（现场约定）", "0.40"),
        ("测试波长(nm)", "否", "默认 1310；1550 业务填 1550", "1550"),
        ("备注", "否", "测试方向/日期/仪表", "OTDR 烧结→白灰窑"),
        ("关联光路编号", "否", "便于对账", "FIB-烧结-SCADA"),
    ]
    for i, row in enumerate(strand_fields, start=70):
        for c, val in enumerate(row, start=1):
            ws.cell(row=i, column=c, value=val).border = Border(
                left=THIN, right=THIN, top=THIN, bottom=THIN
            )

    ws["A77"] = "七、导入后 NetBox 中在哪里看"
    ws["A77"].font = Font(bold=True, size=12)
    ws["A77"].fill = SECTION_FILL
    ws.merge_cells("A77:F77")
    netbox_map = [
        ("缆段每芯光损", "纤缆管理 → 光缆 → [CBL-xxx] → 纤芯 → 实测光损(dB)", "物理缆段档案，每芯一条"),
        ("业务全链路光损", "纤缆管理 → 光路 → 路径 #1 → 详情 → 光学参数", "端到端 OTDR，来自光路录入"),
        ("业务路由/追踪", "纤缆管理 → 光路 → 路径 #1 → 追踪", "可视化全程，含 FrontPort 跳接"),
        ("完整纤芯链", "光路 → 路径 #1 → 详情 → 完整纤芯链", "导入后自动生成，如 轧钢至白灰窑1-1→白灰窑至白灰窑21-1→白灰窑2至烧结1-1"),
        ("端口连接", "设备 → ODF → 连接（FrontPort）", "PATCH/LINK 跳线，来自端口连接表"),
    ]
    style_table_header(ws, 78, ["记录类型", "NetBox 菜单位置", "说明"], [18, 42, 28])
    for i, row in enumerate(netbox_map, start=79):
        for c, val in enumerate(row, start=1):
            ws.cell(row=i, column=c, value=val).border = Border(
                left=THIN, right=THIN, top=THIN, bottom=THIN
            )

    ws["A84"] = "八、导入前检查"
    ws["A84"].font = Font(bold=True, size=12)
    ws["A84"].fill = SECTION_FILL
    ws.merge_cells("A84:F84")
    checks = [
        "1. 光路段：相邻段 B端 与 下一段 A端 的 ODF+纤芯 应能衔接",
        "2. 光路段：同一段内 A/B 纤芯通常相同（表示同一根纤芯两端）",
        "3. 熔接：光路经过的每个「多缆交汇 ODF」都需登记熔接行",
        "4. A/B端设备名必须与 NetBox 完全一致，统一格式 {机房}-{机柜}-ODF-{序号}",
        "5. 缆段编号必须与 NetBox 中 DCIM 缆线 label 完全一致",
        "6. 缆段纤芯光损：缆段+纤芯必须在 NetBox 已存在（先 bootstrap 或导入基础设施）",
        "7. 导入后若提示「路径未完整」，检查熔接是否遗漏或缆段编号错误",
    ]
    for i, text in enumerate(checks, start=85):
        ws[f"A{i}"] = text
        ws.merge_cells(f"A{i}:F{i}")

    ws.freeze_panes = "A17"


def write_data_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("光路录入")
    style_table_header(ws, 1, DATA_HEADERS, DATA_WIDTHS)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 34
    for row in EXAMPLE_ROWS:
        ws.append(row)
    note_row = 2 + len(EXAMPLE_ROWS) + 1
    ws.merge_cells(f"A{note_row}:{LAST_COL}{note_row}")
    ws[f"A{note_row}"] = (
        "相同光路编号=多段。纤芯填 排-芯，端口名(参考)列演示导入结果；端口连接见「端口连接」。详见「字段说明」。"
    )
    ws[f"A{note_row}"].font = NOTE_FONT


def write_port_link_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("端口连接")
    style_table_header(ws, 1, PORT_LINK_HEADERS, PORT_LINK_WIDTHS, SPLICE_HEADER_FILL)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 34
    for row in PORT_LINK_EXAMPLE_ROWS:
        ws.append(row)
    note_row = 2 + len(PORT_LINK_EXAMPLE_ROWS) + 1
    ws.merge_cells(f"A{note_row}:{PORT_LINK_LAST_COL}{note_row}")
    ws[f"A{note_row}"] = (
        "导入时在 NetBox 创建 FrontPort↔FrontPort 连接（同机房 PATCH-* / 跨框 LINK-*）。"
        "光路相邻段换缆时通常自动推断，本表仅补充或覆盖。"
    )
    ws[f"A{note_row}"].font = NOTE_FONT


def write_strand_loss_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("缆段纤芯光损")
    style_table_header(ws, 1, STRAND_LOSS_HEADERS, STRAND_LOSS_WIDTHS, STRAND_LOSS_HEADER_FILL)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 34
    for row in STRAND_LOSS_EXAMPLE_ROWS:
        ws.append(row)
    note_row = 2 + len(STRAND_LOSS_EXAMPLE_ROWS) + 1
    ws.merge_cells(f"A{note_row}:{STRAND_LOSS_LAST_COL}{note_row}")
    ws[f"A{note_row}"] = (
        "每行=一条 CBL 缆段上的一根纤芯 OTDR 实测。导入写入 NetBox 纤芯「实测光损(dB)」。"
        "业务全链路光损仍填在「光路录入」末段。"
    )
    ws[f"A{note_row}"].font = NOTE_FONT


SYSTEM_MAP_ROWS = [
    ("DCIM 站点", "name", "厂区站点登记", "站点名称*", "是", "信通预建", "NetBox → 组织 → 站点"),
    ("DCIM 站点", "slug", "厂区站点登记", "URL标识(slug)*", "是", "信通预建", "英文标识，如 defeng-coking"),
    ("DCIM 站点", "latitude", "厂区站点登记", "纬度", "建议", "信通预建", "OSP 地图定位"),
    ("DCIM 站点", "longitude", "厂区站点登记", "经度", "建议", "信通预建", "OSP 地图定位"),
    ("DCIM 设备", "name", "ODF设备登记", "设备名称*", "是", "信通预建", "格式：{机房}-{机柜}-ODF-{序号}，全表统一"),
    ("DCIM 设备", "site", "ODF设备登记", "站点名称*", "是", "信通预建", "关联站点"),
    ("DCIM 设备", "role", "ODF设备登记", "设备角色", "是", "信通预建", "固定 ODF"),
    ("DCIM 设备", "device_type", "ODF设备登记", "设备型号", "是", "信通预建", "如 ODF-48"),
    ("DCIM 缆线", "label", "厂区缆段登记", "缆段编号*", "是", "信通预建", "CBL-{起点}-{终点}，光路表引用"),
    ("DCIM 缆线", "length", "厂区缆段登记", "缆段长度(m)", "建议", "信通预建", ""),
    ("DCIM 后端口", "name", "厂区缆段登记", "A/B端后端口名(系统)", "自动", "信通/脚本", "RP-{缆段编号}-A/B"),
    ("FMS 光缆类型", "model", "厂区缆段登记", "光纤型号*", "是", "信通预建", "如 G.652D-12"),
    ("FMS 光缆类型", "strand_count", "厂区缆段登记", "芯数*", "是", "信通预建", "与 ODF 端口数一致"),
    ("FMS 光缆", "cable", "厂区缆段登记", "缆段编号*", "是", "信通预建", "绑定 DCIM 缆线"),
    ("DCIM 前置端口", "name", "光路录入", "A/B端纤芯(排-芯)*", "是", "导入生成", "导入后重命名为 {起点}至{终点}{排}-{芯}"),
    ("DCIM 跳线", "label", "端口连接", "—", "自动", "导入生成", "PATCH-* 同机房 / LINK-* 跨框"),
    ("FMS 光路", "cid", "光路录入", "光路编号*", "是", "导入", "业务链路唯一编号"),
    ("FMS 光路", "name", "光路录入", "光路名称", "首段", "导入", ""),
    ("FMS 光路", "description", "光路录入", "业务说明", "否", "导入", "写入光路描述"),
    ("FMS 路径", "actual_loss_db", "光路录入", "实测光损(dB)", "末段", "导入", "端到端 OTDR"),
    ("FMS 路径", "wavelength_nm", "光路录入", "测试波长(nm)", "否", "导入", "1310 或 1550"),
    ("FMS 路径", "fiber_chain", "—", "—", "自动", "导入生成", "完整纤芯链，路径追踪页显示"),
    ("FMS 路径", "fiber_chain_detail", "—", "—", "自动", "导入生成", "逐站纤芯明细"),
    ("FMS 纤芯", "measured_loss_db", "缆段纤芯光损", "实测光损(dB)*", "是", "导入", "每缆每芯 OTDR"),
    ("FMS 纤芯", "test_wavelength_nm", "缆段纤芯光损", "测试波长(nm)", "否", "导入", "自定义字段"),
    ("FMS 熔接条目", "fiber_a/fiber_b", "熔接录入", "A/B端纤芯*", "是", "导入(熔接模式)", "端口模式可改用「端口连接」"),
    ("OSP 光缆", "name", "厂区缆段登记", "缆段编号*", "是", "同步脚本", "sync_osp_strands.py 从 DCIM 同步"),
]


def write_system_map_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("系统字段对照", 1)
    style_table_header(ws, 1, SYSTEM_MAP_HEADERS, SYSTEM_MAP_WIDTHS, PatternFill("solid", fgColor="404040"))
    ws.freeze_panes = "A2"
    for row in SYSTEM_MAP_ROWS:
        ws.append(list(row))
    for r in range(2, 2 + len(SYSTEM_MAP_ROWS)):
        for c in range(1, len(SYSTEM_MAP_HEADERS) + 1):
            ws.cell(row=r, column=c).border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _write_infra_sheet(
    wb: Workbook,
    title: str,
    headers: list[str],
    widths: list[int],
    rows: list[list],
    note: str,
    fill: PatternFill,
) -> None:
    ws = wb.create_sheet(title)
    style_table_header(ws, 1, headers, widths, fill)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 34
    for row in rows:
        ws.append(row)
    last_col = get_column_letter(len(headers))
    note_row = 2 + len(rows) + 1
    ws.merge_cells(f"A{note_row}:{last_col}{note_row}")
    ws[f"A{note_row}"] = note
    ws[f"A{note_row}"].font = NOTE_FONT


def write_site_sheet(wb: Workbook) -> None:
    _write_infra_sheet(
        wb,
        "厂区站点登记",
        SITE_HEADERS,
        SITE_WIDTHS,
        SITE_EXAMPLE_ROWS,
        "信通在 NetBox 预建站点。纬度/经度用于 OSP 离线地图定位。填写后运行 bootstrap_test_infra.py 或手工录入。",
        SITE_HEADER_FILL,
    )


def write_device_sheet(wb: Workbook) -> None:
    _write_infra_sheet(
        wb,
        "ODF设备登记",
        DEVICE_HEADERS,
        DEVICE_WIDTHS,
        DEVICE_EXAMPLE_ROWS,
        "每行一台 ODF。「设备名称*」必须与光路/缆段/端口连接表中 A/B 端设备完全一致。"
        "命名：{机房}-{机柜}-ODF-{序号}。导入前须已在 NetBox 存在。",
        DEVICE_HEADER_FILL,
    )


def write_trunk_sheet(wb: Workbook) -> None:
    _write_infra_sheet(
        wb,
        "厂区缆段登记",
        TRUNK_HEADERS,
        TRUNK_WIDTHS,
        TRUNK_EXAMPLE_ROWS,
        "每行一条厂区主干缆 CBL-*。信通预建 DCIM 缆线 + FMS 光缆 + 后端口。"
        "后端口名列可留空，bootstrap 脚本自动生成 RP-{缆段编号}-A/B。",
        TRUNK_HEADER_FILL,
    )


def write_splice_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("熔接录入")
    style_table_header(ws, 1, SPLICE_HEADERS, SPLICE_WIDTHS, SPLICE_HEADER_FILL)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 34
    for row in SPLICE_EXAMPLE_ROWS:
        ws.append(row)
    note_row = 2 + len(SPLICE_EXAMPLE_ROWS) + 1
    ws.merge_cells(f"A{note_row}:{SPLICE_LAST_COL}{note_row}")
    ws[f"A{note_row}"] = (
        "仅在使用 import_from_excel.py --link-mode splice 时需要。"
        "默认端口模式请用「端口连接」表，本表可留空。"
    )
    ws[f"A{note_row}"].font = NOTE_FONT


def main() -> None:
    from generate_import_sheet import main as export_main

    export_main()


if __name__ == "__main__":
    main()
