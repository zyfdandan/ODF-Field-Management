#!/usr/bin/env python3
"""Parse 公司光缆资产表111.xlsx -> import confirmation workbook."""

from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from naming_rules import (
    asset_room_name,
    assign_room_abbrevs,
    format_cable_label,
    format_odf_device_name,
    infer_route_suffix,
    normalize_cabinet_code,
    port_name_for_cable_fiber,
    room_abbrev,
)

SOURCE_CANDIDATES = [
    Path(__file__).with_name("光纤走向_去重.tsv"),
    Path(__file__).with_name("光纤走向_去重.xlsx"),
    Path(r"C:/Users/Administrator/Downloads/公司光缆资产表111.xlsx"),
    Path(r"C:/Users/Administrator/Downloads/公司光缆资产表.xlsx"),
]
OUTPUT = Path(__file__).with_name("公司光缆资产_导入确认表.xlsx")
CLEAN_XLSX = Path(__file__).with_name("光纤走向_去重.xlsx")
SITE = "得丰焦化"


def resolve_source() -> Path:
    for p in SOURCE_CANDIDATES:
        if p.exists():
            return p
    raise FileNotFoundError(f"找不到源文件，已尝试: {', '.join(str(p) for p in SOURCE_CANDIDATES)}")


def _parse_cabinet(raw: str) -> str:
    raw = (raw or "").strip()
    if not raw:
        return ""
    if re.search(r"[A-Za-z]\d", raw):
        return normalize_cabinet_code(raw)
    return ""


def _parse_cores(raw: str) -> str:
    s = (raw or "").strip()
    m = re.search(r"(\d+)", s)
    return m.group(1) if m else s.replace("芯", "")


class RouteNaming:
    """机房至机房 ODF/缆段命名；同机房对第二条起加「备用」。"""

    def __init__(self, abbrev_registry: dict[str, str]) -> None:
        self._abbrev = abbrev_registry
        self._pair_count: dict[tuple[str, str], int] = defaultdict(int)

    def segment_suffix(self, start: str, end: str, note: str = "") -> str:
        start, end = asset_room_name(start), asset_room_name(end)
        explicit = infer_route_suffix(note)
        if explicit:
            return explicit
        pair = tuple(sorted((start, end)))
        self._pair_count[pair] += 1
        return "" if self._pair_count[pair] == 1 else "备用"

    def cable(self, start: str, end: str, *, suffix: str) -> str:
        return format_cable_label(start, end, suffix=suffix, from_room=start)

    def odf_at(self, local: str, peer: str, *, suffix: str = "") -> str:
        return format_odf_device_name(local, peer, suffix=suffix)


def _segment_key(start: str, end: str, cores: str, length: str) -> tuple[str, str, str, str]:
    return (asset_room_name(start), asset_room_name(end), cores, length)


def dedupe_segments(segments: list[dict[str, str]]) -> tuple[list[dict[str, str]], list[str]]:
    """去掉完全相同的缆段（起终点+芯数+长度），保留首次出现。"""
    seen: set[tuple[str, str, str, str]] = set()
    out: list[dict[str, str]] = []
    notes: list[str] = []
    for s in segments:
        key = _segment_key(s["起点"], s["终点"], s["芯数"], s["长度(米)"])
        if key in seen:
            notes.append(f"已跳过重复: {s['起点']} → {s['终点']} {s['芯数']}芯 {s['长度(米)']}m")
            continue
        seen.add(key)
        out.append(s)
    for i, s in enumerate(out, 1):
        s["序号"] = str(i)
    return out, notes


def _load_route_dataframe(source: Path) -> pd.DataFrame:
    if source.suffix.lower() == ".tsv":
        return pd.read_csv(source, sep="\t")
    return pd.read_excel(source, sheet_name="光纤走向")


def parse_route_segments(source: Path) -> list[dict[str, str]]:
    df = _load_route_dataframe(source)
    cols = [str(c).strip() for c in df.columns]
    if len(cols) >= 4 and cols[0] == "机房位置" and cols[1] == "机房位置":
        pass
    elif len(cols) >= 4 and "机房" in cols[0]:
        df = df.iloc[1:].reset_index(drop=True)
    out: list[dict[str, str]] = []
    for _, row in df.iterrows():
        start = asset_room_name(str(row.iloc[0]))
        end = asset_room_name(str(row.iloc[1]))
        if not start or not end or start.lower() == "nan" or end.lower() == "nan":
            continue
        cores = _parse_cores(str(row.iloc[2]) if len(row) > 2 else "")
        length = row.iloc[3] if len(row) > 3 else ""
        if pd.isna(length):
            len_s = ""
        else:
            try:
                len_s = str(int(float(length)))
            except (TypeError, ValueError):
                len_s = str(length).strip()
        out.append(
            {
                "序号": str(len(out) + 1),
                "起点": start,
                "终点": end,
                "芯数": cores,
                "长度(米)": len_s,
                "备注": "",
            }
        )
    return out


def write_clean_xlsx(segments: list[dict[str, str]]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "光纤走向"
    ws.append(["机房位置", "机房位置", "芯数", "长度"])
    for s in segments:
        cores = s["芯数"]
        core_label = f"{cores}芯" if cores and not str(cores).endswith("芯") else cores
        ws.append([s["起点"], s["终点"], core_label, s["长度(米)"] or None])
    wb.save(CLEAN_XLSX)


def collect_rooms(segments: list[dict[str, str]]) -> list[str]:
    rooms: set[str] = set()
    for s in segments:
        rooms.add(s["起点"])
        rooms.add(s["终点"])
    return sorted(rooms)


def build_base_config(
    reg: RouteNaming, rooms: list[str], segments: list[dict[str, str]], source_name: str
) -> list[list[str]]:
    rows: list[list[str]] = []
    rows.append(["站点", SITE, "defeng-coking", "", f"来自{source_name}"])
    for room in rooms:
        rows.append(["机房", room, SITE, "", ""])
    seen_odf: set[str] = set()
    for s in segments:
        suffix = s.get("路由后缀", "")
        for local, peer in ((s["起点"], s["终点"]), (s["终点"], s["起点"])):
            odf = reg.odf_at(local, peer, suffix=suffix)
            if odf in seen_odf:
                continue
            seen_odf.add(odf)
            rows.append(["ODF", odf, SITE, local, f"对端{peer}"])
    seen_cables: set[str] = set()
    for s in segments:
        label = s["缆段编号(建议)"]
        if not label or label in seen_cables:
            continue
        seen_cables.add(label)
        suffix = s.get("路由后缀", "")
        dev_a = reg.odf_at(s["起点"], s["终点"], suffix=suffix)
        dev_b = reg.odf_at(s["终点"], s["起点"], suffix=suffix)
        note = f"芯数{s['芯数']} 长度{s['长度(米)']}m".strip()
        rows.append(["缆段", label, dev_a, dev_b, note])
    return rows


def build_odf_device_rows(base_rows: list[list[str]], segments: list[dict[str, str]], reg: RouteNaming) -> list[list[str]]:
    odf_cores: dict[str, str] = {}
    for s in segments:
        suffix = s.get("路由后缀", "")
        cores = str(s.get("芯数") or "12").strip() or "12"
        dev_a = reg.odf_at(s["起点"], s["终点"], suffix=suffix)
        dev_b = reg.odf_at(s["终点"], s["起点"], suffix=suffix)
        odf_cores[dev_a] = cores
        odf_cores[dev_b] = cores

    out: list[list[str]] = []
    for row in base_rows:
        if not row or row[0] != "ODF":
            continue
        odf, room = row[1], row[3]
        cores = odf_cores.get(odf, "12")
        out.append([SITE, room, odf, "ODF", f"ODF-{cores}", "", "", "", "", "active", row[4] if len(row) > 4 else ""])
    return out


def write_output(
    source: Path,
    rooms: list[str],
    segments: list[dict[str, str]],
    base_rows: list[list[str]],
    reg: RouteNaming,
    odf_device_rows: list[list[str]],
    abbrev_registry: dict[str, str],
) -> None:
    wb = openpyxl.Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)

    def style_header(ws):
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    room_odfs: dict[str, set[str]] = defaultdict(set)
    for s in segments:
        suffix = s.get("路由后缀", "")
        room_odfs[s["起点"]].add(reg.odf_at(s["起点"], s["终点"], suffix=suffix))
        room_odfs[s["终点"]].add(reg.odf_at(s["终点"], s["起点"], suffix=suffix))

    ws1 = wb.active
    ws1.title = "机房清单"
    h1 = ["序号", "站点", "机房名称", "站点缩写", "关联ODF设备", "备注"]
    ws1.append(h1)
    for i, room in enumerate(rooms, 1):
        odfs = "、".join(sorted(room_odfs.get(room, [])))
        ws1.append([i, SITE, room, abbrev_registry.get(room, room_abbrev(room)), odfs, ""])
    style_header(ws1)
    for col, w in enumerate([6, 12, 24, 12, 48, 16], 1):
        ws1.column_dimensions[get_column_letter(col)].width = w

    ws2 = wb.create_sheet("光缆缆段")
    h2 = ["序号", "起点机房", "终点机房", "缆段编号", "芯数", "长度(米)", "A端ODF", "B端ODF", "A端端口示例", "B端端口示例"]
    ws2.append(h2)
    for s in segments:
        suffix = s.get("路由后缀", "")
        cab = s["缆段编号(建议)"]
        pa = port_name_for_cable_fiber(cab, "1-1", local_room=s["起点"], abbrev_registry=abbrev_registry)
        pb = port_name_for_cable_fiber(cab, "1-1", local_room=s["终点"], abbrev_registry=abbrev_registry)
        ws2.append(
            [
                s["序号"],
                s["起点"],
                s["终点"],
                cab,
                s["芯数"],
                s["长度(米)"],
                reg.odf_at(s["起点"], s["终点"], suffix=suffix),
                reg.odf_at(s["终点"], s["起点"], suffix=suffix),
                pa,
                pb,
            ]
        )
    style_header(ws2)
    for col, w in enumerate([6, 18, 18, 36, 8, 10, 36, 36, 16, 16], 1):
        ws2.column_dimensions[get_column_letter(col)].width = w

    ws3 = wb.create_sheet("基础配置")
    ws3.append(["类型", "名称", "关联A", "关联B", "备注"])
    for row in base_rows:
        ws3.append(row)
    style_header(ws3)

    ws4 = wb.create_sheet("ODF设备登记")
    ws4.append(
        ["站点名称*", "机房/位置*", "设备名称*", "设备角色", "设备型号", "机柜号", "U位", "纬度", "经度", "状态", "备注"]
    )
    for row in odf_device_rows:
        ws4.append(row)
    style_header(ws4)

    ws5 = wb.create_sheet("填写说明")
    for n in [
        ["源文件: " + source.name],
        ["机房：按资产表「机房位置」原样导入 NetBox Location"],
        ["ODF/缆段：{本端机房}至{对端机房}；同机房对第二条缆加「备用」"],
        ["端口（导入后）：{本端缩写}to{对端缩写}{排}-{芯}，如 DFtoXZJ1-1；每排12芯"],
        ["ODF 型号：先读缆段芯数，再建 ODF-{芯数} 设备（如 ODF-96），端口数与芯数一致"],
        ["确认后：python import_asset_infra.py  或合并「基础配置」后 bootstrap_from_sheet.py"],
    ]:
        ws5.append(n)
    ws5.column_dimensions["A"].width = 100

    wb.save(OUTPUT)
    print(f"Wrote {OUTPUT}")
    print(f"  源文件 {source.name}")
    print(f"  机房 {len(rooms)}  缆段 {len(segments)}  ODF {len(odf_device_rows)}")


def main() -> None:
    source = resolve_source()
    segments_raw = parse_route_segments(source)
    if not segments_raw:
        raise SystemExit("光纤走向表无有效缆段")
    before = len(segments_raw)
    segments_raw, skipped = dedupe_segments(segments_raw)
    if skipped:
        print(f"去重: {before} -> {len(segments_raw)} 条（跳过 {len(skipped)} 条完全重复）")
        for line in skipped:
            print(f"  {line}")
    write_clean_xlsx(segments_raw)
    print(f"已写入清洁源表: {CLEAN_XLSX}")
    rooms = collect_rooms(segments_raw)
    abbrev_registry = assign_room_abbrevs(rooms)
    reg = RouteNaming(abbrev_registry)
    segments: list[dict[str, str]] = []
    for s in segments_raw:
        note = s.get("备注", "")
        suffix = reg.segment_suffix(s["起点"], s["终点"], note)
        label = reg.cable(s["起点"], s["终点"], suffix=suffix)
        segments.append({**s, "缆段编号(建议)": label, "路由后缀": suffix})
    base = build_base_config(reg, rooms, segments, source.name)
    odf_rows = build_odf_device_rows(base, segments, reg)
    write_output(source, rooms, segments, base, reg, odf_rows, abbrev_registry)


if __name__ == "__main__":
    main()
