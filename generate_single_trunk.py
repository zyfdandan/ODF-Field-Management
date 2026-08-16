#!/usr/bin/env python3
"""Generate a minimal single-trunk import sheet (基础配置 only)."""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from naming_rules import format_cable_label, format_odf_device_name, port_name_for_cable_fiber, site_slug
from unified_sheet import BASE_HEADERS, SHEET_BASE

ROOT = Path(__file__).resolve().parent
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def main() -> None:
    ap = argparse.ArgumentParser(description="生成单条缆段最小导入表（仅基础配置）")
    ap.add_argument("--room-a", default="焦化老浴池", help="A端机房")
    ap.add_argument("--room-b", default="得丰办公楼梯井", help="B端机房")
    ap.add_argument("--site", default="得丰焦化", help="所属站点")
    ap.add_argument("--cores", type=int, default=4, help="芯数")
    ap.add_argument("--length", default="", help="长度(米)，可留空")
    ap.add_argument("-o", "--output", type=Path, default=None)
    args = ap.parse_args()

    room_a, room_b = args.room_a.strip(), args.room_b.strip()
    site = args.site.strip()
    cable = format_cable_label(room_a, room_b, from_room=room_a)
    odf_a = format_odf_device_name(room_a, room_b)
    odf_b = format_odf_device_name(room_b, room_a)
    note = f"芯数{args.cores}"
    if args.length:
        note += f" 长度{args.length}m"

    out = args.output or ROOT / f"新增缆段_{room_a}至{room_b}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_BASE
    for col, h in enumerate(BASE_HEADERS, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    rows = [
        ["机房", room_a, site, "", ""],
        ["机房", room_b, site, "", ""],
        ["ODF", odf_a, site, room_a, f"对端{room_b}"],
        ["ODF", odf_b, site, room_b, f"对端{room_a}"],
        ["缆段", cable, odf_a, odf_b, note],
    ]
    for row in rows:
        ws.append(row)

    ws2 = wb.create_sheet("说明")
    pa = port_name_for_cable_fiber(cable, "1-1", local_room=room_a)
    pb = port_name_for_cable_fiber(cable, "1-1", local_room=room_b)
    tips = [
        "【只加缆段时】只填「基础配置」这 5 行即可，其它工作表不用管。",
        "",
        f"缆段编号：{cable}",
        f"A端 ODF：{odf_a}（{room_a}）",
        f"B端 ODF：{odf_b}（{room_b}）",
        f"芯数：{args.cores} → 导入后端口 1-1 … 1-{args.cores}",
        f"端口示例：{room_a}侧 {pa}，{room_b}侧 {pb}",
        "",
        "导入命令：",
        f"  python import_asset_infra.py --xlsx {out.name}",
        "",
        "若还要登记业务光路，再填「现场登记」1 行；有 OTDR 再填「光损录入」。",
    ]
    for i, line in enumerate(tips, 1):
        ws2.cell(row=i, column=1, value=line)
    ws2.column_dimensions["A"].width = 80

    wb.save(out)
    print(f"已写入 {out}")
    print(f"  缆段 {cable} · {args.cores}芯 · ODF {odf_a} / {odf_b}")


if __name__ == "__main__":
    main()
