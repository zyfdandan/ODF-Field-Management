#!/usr/bin/env python3
"""Generate and embed QR code images in Excel worksheets."""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

QR_SIZE_PX = 96
QR_ROW_HEIGHT = 72
QR_COL_WIDTH = 14

QR_COLUMN_NAMES = ("二维码", "二维码文件")


def save_qr_png(path: Path, url: str) -> bool:
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        import qrcode

        qrcode.make(url).save(path)
        return True
    except ImportError:
        print("  提示: 安装 qrcode[pil] 后可生成并嵌入二维码 (pip install qrcode[pil])")
        return False
    except OSError as e:
        print(f"  二维码保存失败 {path}: {e}")
        return False


def embed_qr_in_cell(ws: Worksheet, row_idx: int, col_idx: int, png_path: Path) -> bool:
    if not png_path.is_file():
        return False
    try:
        img = Image(str(png_path))
        img.width = QR_SIZE_PX
        img.height = QR_SIZE_PX
        anchor = f"{get_column_letter(col_idx)}{row_idx}"
        ws.add_image(img, anchor)
        current_h = ws.row_dimensions[row_idx].height or 0
        if current_h < QR_ROW_HEIGHT:
            ws.row_dimensions[row_idx].height = QR_ROW_HEIGHT
        col_letter = get_column_letter(col_idx)
        current_w = ws.column_dimensions[col_letter].width or 0
        if current_w < QR_COL_WIDTH:
            ws.column_dimensions[col_letter].width = QR_COL_WIDTH
        return True
    except OSError as e:
        print(f"  二维码嵌入失败 row={row_idx}: {e}")
        return False


def find_qr_column(headers: list[str]) -> int | None:
    for name in QR_COLUMN_NAMES:
        if name in headers:
            return headers.index(name) + 1
    return None


def write_qr_back_to_excel(
    wb: Workbook,
    excel_path: Path,
    *,
    cid_col: int,
    col_url: int,
    col_qr: int,
    results: list[dict[str, str]],
    sheet_name: str = "光路录入",
) -> int:
    """Write trace URLs and embed QR images (one per 光路编号, first segment row)."""
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
    ws._images = []
    by_cid = {r["cid"]: r for r in results}
    seen: set[str] = set()
    embedded = 0

    for row_idx in range(2, ws.max_row + 1):
        raw = ws.cell(row=row_idx, column=cid_col).value
        if not raw:
            continue
        cid = str(raw).strip()
        if cid not in by_cid:
            continue
        info = by_cid[cid]
        ws.cell(row=row_idx, column=col_url, value=info["url"])
        if cid in seen:
            continue
        seen.add(cid)
        ws.cell(row=row_idx, column=col_qr, value=None)
        abs_path = (info.get("qr_abs_path") or "").strip()
        if not abs_path:
            continue
        png = Path(abs_path)
        if embed_qr_in_cell(ws, row_idx, col_qr, png):
            embedded += 1

    try:
        wb.save(excel_path)
    except PermissionError:
        alt = excel_path.with_name(excel_path.stem + "_导入回写.xlsx")
        wb.save(alt)
        print(f"Excel 被占用，已另存为: {alt}")
    return embedded


def embed_from_manifest(
    excel_path: Path,
    manifest_path: Path,
    *,
    sheet_name: str = "光路录入",
    cid_header: str = "光路编号",
    url_header: str = "扫码链接",
) -> int:
    """Embed QR PNGs listed in qr_manifest.json without re-importing."""
    import json

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    wb = openpyxl.load_workbook(excel_path)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
    headers = [str(c.value or "").replace("*", "") for c in ws[1]]
    cid_col = headers.index(cid_header.replace("*", "")) + 1
    col_url = headers.index(url_header) + 1
    col_qr = find_qr_column(headers)
    if not col_qr:
        raise ValueError(f"缺少列: {' / '.join(QR_COLUMN_NAMES)}")

    base = excel_path.parent
    results = []
    for item in manifest:
        rel = item.get("qr_file", "")
        png = base / rel if rel else base / "qr_labels" / f"{item['cid']}.png"
        results.append(
            {
                "cid": item["cid"],
                "url": item.get("url", ""),
                "qr_abs_path": str(png.resolve()),
            }
        )
    return write_qr_back_to_excel(
        wb,
        excel_path,
        cid_col=cid_col,
        col_url=col_url,
        col_qr=col_qr,
        results=results,
        sheet_name=sheet_name,
    )
