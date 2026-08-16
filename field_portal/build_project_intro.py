#!/usr/bin/env python3
"""Generate concise ODF Field Portal intro for leadership briefing."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "ODF现场管理系统-领导汇报.docx"


def set_run_font(run, size=12, bold=False, color=None, east="宋体", ascii_font="Calibri"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = ascii_font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), east)
    if color:
        run.font.color.rgb = color


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=15 if level == 1 else 13, bold=True, east="黑体")


def p(doc, text, *, indent=True, bold=False, size=12):
    para = doc.add_paragraph()
    pf = para.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(4)
    if indent:
        pf.first_line_indent = Cm(0.74)
    run = para.add_run(text)
    set_run_font(run, size=size, bold=bold, east="宋体")


def bullets(doc, items):
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        para.paragraph_format.space_after = Pt(2)
        run = para.add_run(item)
        set_run_font(run, size=12, east="宋体")


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(2.4))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    r = title.add_run("ODF 现场管理系统")
    set_run_font(r, size=20, bold=True, east="黑体")

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(16)
    r = sub.add_run("项目汇报材料（功能 · 优势 · 架构）")
    set_run_font(r, size=12, east="宋体", color=RGBColor(0x44, 0x44, 0x44))

    # 一、概述
    h(doc, "一、项目概述")
    p(
        doc,
        "本系统基于 NetBox 光网络资源库，面向机房 ODF 配线现场运维，"
        "支持通过 VPN 安全访问内网管理页面，在线查看机房面板、登记业务、调整跳纤换芯，"
        "操作结果实时写入 NetBox，实现光路数据统一、可查、可追溯。",
    )

    # 二、使用方式
    h(doc, "二、使用方式")
    bullets(
        doc,
        [
            "接入单位 VPN，进入内网后浏览器打开现场管理系统地址。",
            "按「厂区 → 机房 → ODF 框」浏览端口占用与链路，直接在页面完成登记与修改。",
            "无需现场扫码贴码；授权人员凭 VPN 与账号权限即可开展日常维护。",
        ],
    )

    # 三、主要功能
    h(doc, "三、主要功能")
    bullets(
        doc,
        [
            "机房面板：一屏查看各 ODF 框端口占用、业务归属与光损信息。",
            "光路查询：展示端到端链路轨迹，便于开通核对与故障定位。",
            "业务登记：在线登记光路、跳纤及对端设备连接，数据同步至 NetBox。",
            "跳纤换芯：支持选择空闲纤芯并级联调整，提交前可预览变更结果。",
            "光损维护：按端口补录、更新光损，便于质量跟踪。",
            "顺序调整：可按现场机柜实际排布调整 ODF 显示顺序。",
        ],
    )

    # 四、主要优势
    h(doc, "四、主要优势")
    bullets(
        doc,
        [
            "数据统一：以 NetBox 为唯一权威数据源，减少台账不一致。",
            "操作在线：VPN 内访问即可改，替代纸质记录与反复 Excel 对账。",
            "安全可控：仅内网开放，密钥与写权限留在服务端，降低外泄风险。",
            "响应更快：缓存与分步加载，大机房场景下页面可用、可改。",
            "改线可控：换芯支持预览确认，降低误操作影响。",
            "建设轻量：部署简单、运维成本低，可与既有 NetBox 体系衔接推广。",
        ],
    )

    # 五、系统架构
    h(doc, "五、系统架构")
    p(
        doc,
        "访问路径：终端经 VPN → 内网门户/管理页 → 现场管理服务（Field API）→ NetBox。"
        "前端为轻量 Web 页面；后端为 Python 现场服务，负责业务校验、缓存加速及对 NetBox 的读写；"
        "NetBox 保存厂区、机房、ODF、缆线、跳纤与光路等资源数据。",
    )
    p(
        doc,
        "分层示意：展示层（Web）／业务层（现场 API）／数据层（NetBox）。"
        "对外不暴露数据库直连，变更均经业务接口落库，便于审计与权限控制。",
    )

    # 六、请示话术
    h(doc, "六、请示话术（供汇报选用）")
    p(
        doc,
        "领导好。为提升光配线现场管理效率、统一光路台账，我们建设了 ODF 现场管理系统。"
        "系统依托现有 NetBox 资源库，运维人员通过 VPN 进入内网即可查看机房面板、登记业务、办理跳纤换芯，"
        "修改结果实时回写，避免多头台账和事后补录。系统部署轻量、访问限定内网，安全可控。"
        "现提请领导原则同意在相关机房推广使用，并明确由网络/信息化部门负责日常运维与权限管理。"
        "妥否，请批示。",
        indent=True,
    )

    p(
        doc,
        "（简版）提请同意启用 ODF 现场管理系统：VPN 内访问、在线维护 ODF/光路数据并同步 NetBox，"
        "以统一台账、规范改线、提升开通与排障效率。请领导审示。",
        indent=True,
    )

    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    foot.paragraph_format.space_before = Pt(18)
    r = foot.add_run("汇报材料 · 简洁版")
    set_run_font(r, size=9, east="宋体", color=RGBColor(0x99, 0x99, 0x99))

    doc.save(str(OUT))
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
