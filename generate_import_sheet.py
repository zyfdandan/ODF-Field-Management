#!/usr/bin/env python3
"""Export NetBox inventory into unified 现场登记.xlsx (current naming rules)."""

from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from bootstrap_room_devices import paginate
from import_from_excel import NetBoxClient
from naming_rules import (
    GUIDE_NAMING_LINES,
    NAMING_RULES,
    assign_room_abbrevs,
    mirror_cable_label,
    parse_route_label,
    port_name_for_cable_fiber,
    room_abbrev,
    site_slug,
)
from unified_sheet import BASE_HEADERS, BUSINESS_HEADERS, LOSS_HEADERS, SHEET_BASE, SHEET_BUSINESS, SHEET_LOSS

ROOT = Path(__file__).resolve().parent
DEFAULT_CONFIG = ROOT / "netbox_config.json"
DEFAULT_OUTPUT = ROOT / "现场登记.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SECTION_FILL = PatternFill("solid", fgColor="D9E2F3")
NOTE_FONT = Font(color="666666", italic=True)
THIN = openpyxl.styles.Border(
    left=openpyxl.styles.Side(style="thin", color="CCCCCC"),
    right=openpyxl.styles.Side(style="thin", color="CCCCCC"),
    top=openpyxl.styles.Side(style="thin", color="CCCCCC"),
    bottom=openpyxl.styles.Side(style="thin", color="CCCCCC"),
)


def _style_header(ws, headers: list[str], widths: list[int] | None = None, fill: PatternFill = HEADER_FILL) -> None:
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = fill
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if widths and col <= len(widths):
            ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 32


def _is_route_odf(name: str) -> bool:
    return bool(name) and "至" in name and not name.upper().startswith(("PATCH-", "LINK-", "CBL-", "TMP-"))


def _is_trunk_label(label: str) -> bool:
    label = (label or "").strip()
    if not label or label.upper().startswith(("PATCH-", "LINK-", "TMP-")):
        return False
    return "至" in label or label.upper().startswith("CBL-")


def _device_room(dev: dict[str, Any]) -> str:
    loc = dev.get("location") or {}
    if isinstance(loc, dict) and loc.get("name"):
        return str(loc["name"])
    return ""


def _device_site(dev: dict[str, Any]) -> str:
    site = dev.get("site") or {}
    if isinstance(site, dict):
        return str(site.get("name") or "")
    return ""


def _resolve_trunk_endpoints(
    client: NetBoxClient,
    label: str,
    devices_by_name: dict[str, dict[str, Any]],
) -> tuple[str, str, str, str]:
    """Return dev_a, dev_b, room_a, room_b for a trunk cable."""
    label = (label or "").strip()
    parsed = parse_route_label(label)
    ra, rb = parsed.get("room_a", ""), parsed.get("room_b", "")
    if ra and rb:
        dev_a = f"{ra}至{rb}" if f"{ra}至{rb}" in devices_by_name else ""
        dev_b = f"{rb}至{ra}" if f"{rb}至{ra}" in devices_by_name else ""
        if dev_a and dev_b:
            return dev_a, dev_b, ra, rb

    hits = client.request("GET", "/dcim/cables/", params={"label": label, "limit": 1})
    if not hits.get("results"):
        return "", "", ra, rb
    detail = client.request("GET", f"/dcim/cables/{hits['results'][0]['id']}/")
    dev_names: list[str] = []
    for side in ("a_terminations", "b_terminations"):
        for term in detail.get(side) or []:
            if term.get("object_type") != "dcim.rearport":
                continue
            rp = client.request("GET", f"/dcim/rear-ports/{term['object_id']}/")
            dev = rp.get("device") or {}
            name = dev.get("name") if isinstance(dev, dict) else ""
            if name and name not in dev_names:
                dev_names.append(name)
    if len(dev_names) >= 2:
        return dev_names[0], dev_names[1], _device_room(devices_by_name.get(dev_names[0], {})), _device_room(
            devices_by_name.get(dev_names[1], {})
        )
    if len(dev_names) == 1 and ra and rb:
        only = dev_names[0]
        other = f"{rb}至{ra}" if only.startswith(ra) else f"{ra}至{rb}"
        if only.startswith(ra):
            return only, other if other in devices_by_name else "", ra, rb
        return other if other in devices_by_name else "", only, ra, rb
    return "", "", ra, rb


def _load_fms_cables(client: NetBoxClient) -> dict[str, dict[str, Any]]:
    out: dict[str, dict[str, Any]] = {}
    try:
        for fc in paginate(client, "/plugins/fms/fiber-cables/"):
            cab = fc.get("cable") or {}
            lbl = cab.get("label") if isinstance(cab, dict) else ""
            if lbl:
                out[lbl] = fc
    except Exception:
        pass
    return out


def _load_dcim_cables(client: NetBoxClient) -> dict[str, dict[str, Any]]:
    out: dict[str, dict[str, Any]] = {}
    for cab in paginate(client, "/dcim/cables/"):
        label = (cab.get("label") or "").strip()
        if label:
            out[label] = cab
    return out


def _fms_strand_count(fms_cables: dict[str, dict[str, Any]], cable_label: str) -> str:
    fc = fms_cables.get(cable_label) or {}
    sc = fc.get("strand_count")
    if not sc:
        typ = fc.get("type") or {}
        sc = typ.get("strand_count") if isinstance(typ, dict) else None
    if sc:
        try:
            return str(int(sc))
        except (TypeError, ValueError):
            return str(sc)
    return "12"


def _cable_length(dcim_cables: dict[str, dict[str, Any]], cable_label: str) -> str:
    cab = dcim_cables.get(cable_label) or {}
    length = cab.get("length")
    if length is None:
        return ""
    try:
        return str(int(float(length)))
    except (TypeError, ValueError):
        return str(length)


def _collect_inventory(client: NetBoxClient) -> dict[str, Any]:
    sites = {s["name"]: s for s in paginate(client, "/dcim/sites/")}
    locations: dict[tuple[str, str], dict[str, Any]] = {}
    for loc in paginate(client, "/dcim/locations/"):
        site = (loc.get("site") or {}).get("name", "")
        locations[(site, loc.get("name", ""))] = loc

    devices_by_name: dict[str, dict[str, Any]] = {}
    odf_devices: list[dict[str, Any]] = []
    for dev in paginate(client, "/dcim/devices/"):
        name = (dev.get("name") or "").strip()
        if not name:
            continue
        devices_by_name[name] = dev
        if _is_route_odf(name):
            odf_devices.append(dev)

    dcim_cables = _load_dcim_cables(client)
    fms_cables = _load_fms_cables(client)

    trunks: list[dict[str, str]] = []
    seen_labels: set[str] = set()
    for label, cab in dcim_cables.items():
        if not _is_trunk_label(label) or label in seen_labels:
            continue
        seen_labels.add(label)
        dev_a, dev_b, room_a, room_b = _resolve_trunk_endpoints(client, label, devices_by_name)
        if not dev_a or not dev_b:
            continue
        strands = _fms_strand_count(fms_cables, label)
        length = _cable_length(dcim_cables, label)
        trunks.append(
            {
                "label": label,
                "dev_a": dev_a,
                "dev_b": dev_b,
                "room_a": room_a or _device_room(devices_by_name.get(dev_a, {})),
                "room_b": room_b or _device_room(devices_by_name.get(dev_b, {})),
                "strands": strands,
                "length": length,
                "site_a": _device_site(devices_by_name.get(dev_a, {})),
                "site_b": _device_site(devices_by_name.get(dev_b, {})),
            }
        )
    trunks.sort(key=lambda x: x["label"])

    return {
        "sites": sites,
        "locations": locations,
        "devices_by_name": devices_by_name,
        "odf_devices": sorted(odf_devices, key=lambda d: d.get("name", "")),
        "trunks": trunks,
        "fms_cables": fms_cables,
    }


def _build_base_rows(inv: dict[str, Any]) -> list[list[str]]:
    rows: list[list[str]] = []
    site_names = sorted(inv["sites"].keys())
    if not site_names:
        rows.append(["站点", "得丰焦化", site_slug("得丰焦化"), "", "NetBox 导出"])
    for site in site_names:
        slug = inv["sites"][site].get("slug") or site_slug(site)
        rows.append(["站点", site, slug, "", "NetBox 导出"])

    loc_seen: set[tuple[str, str]] = set()
    for dev in inv["odf_devices"]:
        site = _device_site(dev)
        room = _device_room(dev)
        if site and room and (site, room) not in loc_seen:
            loc_seen.add((site, room))
            rows.append(["机房", room, site, "", "NetBox Location"])

    for dev in inv["odf_devices"]:
        site = _device_site(dev)
        room = _device_room(dev)
        rows.append(["ODF", dev["name"], site, room, "NetBox 设备"])

    for trunk in inv["trunks"]:
        note = f"芯数{trunk['strands']}"
        if trunk.get("length"):
            note += f" 长度{trunk['length']}m"
        rows.append(["缆段", trunk["label"], trunk["dev_a"], trunk["dev_b"], note])
    return rows


def _build_trunk_detail_rows(inv: dict[str, Any], abbrev_registry: dict[str, str]) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for i, trunk in enumerate(inv["trunks"], start=1):
        label = trunk["label"]
        pa = port_name_for_cable_fiber(label, "1-1", local_room=trunk["room_a"], abbrev_registry=abbrev_registry)
        pb = port_name_for_cable_fiber(label, "1-1", local_room=trunk["room_b"], abbrev_registry=abbrev_registry)
        rows.append(
            [
                i,
                trunk["room_a"],
                trunk["room_b"],
                label,
                trunk["strands"],
                trunk.get("length") or "",
                trunk["dev_a"],
                trunk["dev_b"],
                pa,
                pb,
            ]
        )
    return rows


def _build_room_rows(inv: dict[str, Any], abbrev_registry: dict[str, str]) -> list[list[Any]]:
    room_odfs: dict[tuple[str, str], set[str]] = defaultdict(set)
    room_site: dict[tuple[str, str], str] = {}
    for dev in inv["odf_devices"]:
        site = _device_site(dev)
        room = _device_room(dev)
        if not room:
            continue
        key = (site, room)
        room_odfs[key].add(dev["name"])
        room_site[key] = site

    rows: list[list[Any]] = []
    for i, ((site, room), odfs) in enumerate(sorted(room_odfs.items(), key=lambda x: (x[0][0], x[0][1])), start=1):
        rows.append([i, site, room, abbrev_registry.get(room, room_abbrev(room, registry=abbrev_registry)), "、".join(sorted(odfs)), ""])
    return rows


def _export_business_segments(client: NetBoxClient, inv: dict[str, Any]) -> list[list[str]]:
    rows: list[list[str]] = []
    try:
        circuits = paginate(client, "/plugins/fms/fiber-circuits/")
    except Exception:
        return rows

    for circ in circuits:
        cid = (circ.get("cid") or "").strip()
        if not cid:
            continue
        name = (circ.get("name") or "").strip()
        desc = (circ.get("description") or "").strip()
        paths = paginate(client, "/plugins/fms/fiber-circuit-paths/", {"circuit_id": circ["id"]})
        for path in paths:
            hops = path.get("hops") or []
            for hop in hops:
                if hop.get("type") != "cable":
                    continue
                cable = (hop.get("cable") or {}).get("label") if isinstance(hop.get("cable"), dict) else ""
                dev_a = (hop.get("device_a") or {}).get("name", "") if isinstance(hop.get("device_a"), dict) else ""
                dev_b = (hop.get("device_b") or {}).get("name", "") if isinstance(hop.get("device_b"), dict) else ""
                port_a = hop.get("fiber_a") or hop.get("port_a") or "1-1"
                port_b = hop.get("fiber_b") or hop.get("port_b") or port_a
                if not dev_a or not dev_b:
                    continue
                da = inv["devices_by_name"].get(dev_a, {})
                site = _device_site(da)
                room = _device_room(da)
                peer = dev_b
                loss = path.get("actual_loss_db")
                wl = path.get("wavelength_nm") or ""
                label_pos = f"{dev_b} {port_b}" if dev_b else ""
                rows.append(
                    [
                        site,
                        room,
                        cable,
                        cid,
                        dev_a,
                        peer,
                        name,
                        str(port_a),
                        str(port_b),
                        desc,
                        "",
                        label_pos,
                        f"路径#{path.get('id', '')}",
                        "",
                        "",
                    ]
                )
                if loss not in (None, ""):
                    rows[-1][9] = desc  # keep service in 业务说明
    return rows


def _export_strand_losses(client: NetBoxClient, fms_cables: dict[str, dict[str, Any]]) -> list[list[str]]:
    rows: list[list[str]] = []
    label_by_fc_id = {fc["id"]: lbl for lbl, fc in fms_cables.items()}
    try:
        for strand in paginate(client, "/plugins/fms/fiber-strands/"):
            loss = strand.get("measured_loss_db")
            if loss in (None, ""):
                continue
            fc_id = strand.get("cable")
            if isinstance(fc_id, dict):
                fc_id = fc_id.get("id")
            cable_label = label_by_fc_id.get(int(fc_id) if fc_id else 0, "")
            if not cable_label:
                cab = strand.get("cable") or {}
                if isinstance(cab, dict):
                    inner = cab.get("cable") or {}
                    cable_label = inner.get("label", "") if isinstance(inner, dict) else ""
            pos = strand.get("position") or 0
            row_no = (int(pos) - 1) // 12 + 1 if int(pos) > 0 else 1
            col_no = (int(pos) - 1) % 12 + 1 if int(pos) > 0 else 1
            fiber = f"{row_no}-{col_no}"
            wl = strand.get("test_wavelength_nm") or strand.get("wavelength_nm") or "1310"
            rows.append(["", "", fiber, cable_label, str(loss), str(wl), "NetBox 导出"])
    except Exception:
        pass
    return rows


def _write_naming_sheet(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("命名规范", 0)
    ws["A1"] = "现场登记命名规范（与 naming_rules.py 同步）"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:D1")
    row = 3
    ws.cell(row=row, column=1, value="对象").fill = HEADER_FILL
    ws.cell(row=row, column=1).font = HEADER_FONT
    for col, title in enumerate(["格式", "示例", "说明"], start=2):
        cell = ws.cell(row=row, column=col, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    row += 1
    for obj, fmt, example, note in NAMING_RULES:
        ws.cell(row=row, column=1, value=obj)
        ws.cell(row=row, column=2, value=fmt)
        ws.cell(row=row, column=3, value=example)
        ws.cell(row=row, column=4, value=note)
        for c in range(1, 5):
            ws.cell(row=row, column=c).border = THIN
        row += 1
    row += 1
    for line in GUIDE_NAMING_LINES:
        if line:
            ws.cell(row=row, column=1, value=line)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            row += 1
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 36
    ws.column_dimensions["D"].width = 48


def _write_guide_sheet(wb: openpyxl.Workbook, stats: dict[str, int]) -> None:
    ws = wb.create_sheet("填写说明")
    lines = [
        ["本表由 generate_import_sheet.py 从 NetBox 导出并套用当前命名规则生成"],
        [f"站点 {stats['sites']} · 机房 {stats['rooms']} · ODF {stats['odfs']} · 缆段 {stats['trunks']}"],
        [""],
        ["工作表说明："],
        ["· 基础配置 — 站点/机房/ODF/缆段，可 bootstrap_from_sheet.py 或 import_asset_infra.py 同步"],
        ["· 光缆缆段 — 缆段清单与端口命名示例（{本端缩写}to{对端缩写}{排}-{芯}）"],
        ["· 现场登记 — 业务光路分段；相同光路编号=一条链路"],
        ["· 光损录入 — 按 ODF框+纤芯 或 缆段+纤芯 填写 OTDR"],
        ["· 命名规范 — 与系统 import/field portal 一致"],
        [""],
        ["导入命令："],
        ["python import_from_excel.py --excel 现场登记.xlsx"],
        ["python import_asset_infra.py   # 仅基础设施"],
    ]
    for i, line in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=line[0])
    ws.column_dimensions["A"].width = 100


def generate_workbook(client: NetBoxClient, output: Path) -> dict[str, int]:
    inv = _collect_inventory(client)
    rooms = sorted({(_device_site(d), _device_room(d)) for d in inv["odf_devices"] if _device_room(d)})
    abbrev_registry = assign_room_abbrevs([room for _, room in rooms])

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    _write_naming_sheet(wb)

    ws_base = wb.create_sheet(SHEET_BASE)
    _style_header(ws_base, BASE_HEADERS, [10, 40, 18, 18, 36])
    for row in _build_base_rows(inv):
        ws_base.append(row)

    ws_trunk = wb.create_sheet("光缆缆段")
    trunk_headers = ["序号", "起点机房", "终点机房", "缆段编号", "芯数", "长度(米)", "A端ODF", "B端ODF", "A端端口示例", "B端端口示例"]
    _style_header(ws_trunk, trunk_headers, [6, 20, 20, 36, 8, 10, 36, 36, 18, 18])
    for row in _build_trunk_detail_rows(inv, abbrev_registry):
        ws_trunk.append(row)

    ws_room = wb.create_sheet("机房清单")
    room_headers = ["序号", "站点", "机房名称", "站点缩写", "关联ODF设备", "备注"]
    _style_header(ws_room, room_headers, [6, 14, 24, 12, 48, 16])
    for row in _build_room_rows(inv, abbrev_registry):
        ws_room.append(row)

    ws_loss = wb.create_sheet(SHEET_LOSS)
    _style_header(ws_loss, LOSS_HEADERS, [16, 36, 14, 24, 12, 12, 24])
    for row in _export_strand_losses(client, inv["fms_cables"]):
        ws_loss.append(row)

    ws_biz = wb.create_sheet(SHEET_BUSINESS)
    _style_header(ws_biz, BUSINESS_HEADERS, [14, 18, 28, 16, 36, 36, 18, 14, 14, 20, 14, 28, 20, 40, 12])
    biz_rows = _export_business_segments(client, inv)
    for row in biz_rows:
        ws_biz.append(row)

    stats = {
        "sites": len(inv["sites"]),
        "rooms": len(rooms),
        "odfs": len(inv["odf_devices"]),
        "trunks": len(inv["trunks"]),
        "business_rows": len(biz_rows),
        "loss_rows": ws_loss.max_row - 1,
    }
    _write_guide_sheet(wb, stats)

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return stats


def main() -> None:
    parser = argparse.ArgumentParser(description="从 NetBox 导出统一格式现场登记导入表")
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("-o", "--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    cfg = json.loads(args.config.read_text(encoding="utf-8"))
    client = NetBoxClient(cfg["base_url"], cfg["token"], cfg.get("verify_ssl", True))
    stats = generate_workbook(client, args.output)
    print(f"已写入 {args.output}")
    print(
        f"  站点 {stats['sites']} · 机房 {stats['rooms']} · ODF {stats['odfs']} · "
        f"缆段 {stats['trunks']} · 业务行 {stats['business_rows']} · 光损行 {stats['loss_rows']}"
    )


if __name__ == "__main__":
    main()
