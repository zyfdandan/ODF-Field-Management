#!/usr/bin/env python3
"""Parse 现场登记.xlsx — 基础配置 + 光损录入 + 现场登记(业务)."""

from __future__ import annotations

from typing import Any, Protocol

import openpyxl

from naming_rules import site_slug

SHEET_BUSINESS = "现场登记"
SHEET_BASE = "基础配置"
SHEET_LOSS = "光损录入"

BUSINESS_HEADERS = [
    "站点名称",
    "机房/位置",
    "缆段编号",
    "光路编号",
    "所在机房ODF",
    "对端ODF",
    "光路名称",
    "A端纤芯(排-芯)",
    "B端纤芯(排-芯)",
    "业务说明",
    "IP地址",
    "标签位置",
    "备注",
    "扫码链接",
    "二维码",
]

BASE_HEADERS = ["类型", "名称", "关联A", "关联B", "备注"]

LOSS_HEADERS = ["光路编号", "ODF框", "纤芯(排-芯)", "缆段编号", "光损(dB)", "波长(nm)", "备注"]

HEADER_ALIASES = {
    "所在机房 ODF": "所在机房ODF",
    "所在机房ODF*": "所在机房ODF",
    "对端 ODF": "对端ODF",
    "类型*": "类型",
    "名称*": "名称",
    "光损(dB)*": "光损(dB)",
    "A端设备": "所在机房ODF",
    "B端设备": "对端ODF",
}

LOSS_HEADER_ALIASES = {
    **HEADER_ALIASES,
    "ODF框*": "ODF框",
    "ODF设备": "ODF框",
    "ODF设备名": "ODF框",
    "所在机房ODF": "ODF框",
    "所在机房 ODF": "ODF框",
}


def _loss_header_index(ws, required: list[str] | None = None) -> dict[str, int]:
    idx: dict[str, int] = {}
    for i, h in enumerate(_headers(ws)):
        key = LOSS_HEADER_ALIASES.get(h, h).replace("*", "")
        idx[key] = i
    if required:
        for col in required:
            base = col.replace("*", "")
            if base not in idx:
                raise ValueError(f"工作表「{ws.title}」缺少列: {col}")
    return idx


def normalize_cell(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, float):
        if val == int(val):
            return str(int(val))
        return format(val, "g")
    return str(val).strip()


def _headers(ws) -> list[str]:
    return [str(c.value or "").replace("*", "") for c in ws[1]]


def _header_index(ws, required: list[str] | None = None) -> dict[str, int]:
    idx: dict[str, int] = {}
    for i, h in enumerate(_headers(ws)):
        key = HEADER_ALIASES.get(h, h).replace("*", "")
        idx[key] = i
    if required:
        for col in required:
            base = col.replace("*", "")
            if base not in idx:
                raise ValueError(f"工作表「{ws.title}」缺少列: {col}")
    return idx


def _cell(row: tuple, idx: dict[str, int], name: str) -> str:
    key = name.replace("*", "")
    if key not in idx:
        return ""
    return normalize_cell(row[idx[key]])


def workbook_format(wb: openpyxl.Workbook) -> str | None:
    """Return 'business_v2', 'legacy_type', or None."""
    if SHEET_BUSINESS not in wb.sheetnames:
        return None
    ws = wb[SHEET_BUSINESS]
    if ws.max_row < 1:
        return None
    headers = set(_headers(ws))
    if "所在机房ODF" in headers or "所在机房 ODF" in headers:
        return "business_v2"
    if "记录类型" in headers:
        return "legacy_type"
    if "光路编号" in headers and ("A端ODF" in headers or "所在机房ODF" in headers):
        return "business_v2"
    return None


def is_unified_workbook(wb: openpyxl.Workbook) -> bool:
    return workbook_format(wb) is not None


def parse_base_sheet(wb: openpyxl.Workbook) -> dict[str, list[dict[str, str]]]:
    out: dict[str, list] = {"sites": [], "locations": [], "devices": [], "trunks": []}
    if SHEET_BASE not in wb.sheetnames:
        return out
    ws = wb[SHEET_BASE]
    idx = _header_index(ws, ["类型", "名称"])
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        kind = _cell(row, idx, "类型").strip()
        name = _cell(row, idx, "名称")
        if not kind or not name:
            continue
        rel_a = _cell(row, idx, "关联A")
        rel_b = _cell(row, idx, "关联B")
        note = _cell(row, idx, "备注")
        if kind in ("站点", "site"):
            out["sites"].append({"name": name, "slug": rel_a or site_slug(name), "note": note})
        elif kind in ("机房", "location"):
            out["locations"].append({"name": name, "site": rel_a, "note": note})
        elif kind.upper() == "ODF" or kind == "设备":
            out["devices"].append(
                {"name": name, "site": rel_a, "location": rel_b or rel_a, "note": note}
            )
        elif kind in ("缆段", "cable"):
            if not rel_a or not rel_b:
                continue
            out["trunks"].append({"label": name, "dev_a": rel_a, "dev_b": rel_b, "note": note})
    return out


def resolve_cable_for_odf_loss(
    segments: list[dict[str, str]],
    odf: str,
    port: str,
    *,
    cid: str = "",
    cable_hint: str = "",
) -> tuple[str | None, str | None]:
    """Map ODF框 + 纤芯 to 缆段编号 using business segments."""
    odf = (odf or "").strip()
    port = (port or "").strip()
    cable_hint = (cable_hint or "").strip()
    cid = (cid or "").strip()
    if not odf or not port:
        return None, "缺少 ODF框 或 纤芯"

    pool = segments
    if cid:
        pool = [s for s in segments if s.get("cid") == cid]
        if not pool:
            return None, f"光路 {cid} 无业务段，无法解析 ODF框"

    matches: list[tuple[str, str]] = []
    for seg in pool:
        cab = (seg.get("cable") or "").strip()
        if not cab:
            continue
        if seg.get("dev_a") == odf and seg.get("port_a") == port:
            matches.append((cab, "A"))
        if seg.get("dev_b") == odf and seg.get("port_b") == port:
            matches.append((cab, "B"))

    if not matches:
        return None, f"ODF框 {odf} 纤芯 {port} 在业务段中未找到"

    cables = sorted({m[0] for m in matches})
    if cable_hint:
        if cable_hint not in cables:
            return None, f"缆段 {cable_hint} 与 {odf}:{port} 不匹配（可选: {', '.join(cables)}）"
        return cable_hint, None
    if len(cables) == 1:
        return cables[0], None
    return None, f"{odf} 纤芯 {port} 对应多条缆段（{', '.join(cables)}），请填写「缆段编号」列"


def parse_loss_sheet(
    wb: openpyxl.Workbook,
    segments: list[dict[str, str]] | None = None,
) -> tuple[list[dict[str, str]], dict[str, dict[str, str]], list[str]]:
    """Return (strand_losses, circuit_losses_by_cid, warnings)."""
    strand_losses: list[dict[str, str]] = []
    circuit_losses: dict[str, dict[str, str]] = {}
    warnings: list[str] = []
    segments = segments or []
    if SHEET_LOSS not in wb.sheetnames:
        return strand_losses, circuit_losses, warnings
    ws = wb[SHEET_LOSS]
    idx = _loss_header_index(ws, ["光损(dB)"])
    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row:
            continue
        loss = _cell(row, idx, "光损(dB)")
        if not loss:
            continue
        cid = _cell(row, idx, "光路编号")
        odf = _cell(row, idx, "ODF框")
        cable = _cell(row, idx, "缆段编号")
        port = _cell(row, idx, "纤芯(排-芯)")
        wl = _cell(row, idx, "波长(nm)") or "1310"
        note = _cell(row, idx, "备注")

        if odf and port:
            resolved, err = resolve_cable_for_odf_loss(
                segments, odf, port, cid=cid, cable_hint=cable
            )
            if err:
                warnings.append(f"光损录入 第{row_no}行: {err}")
                continue
            cable = resolved or cable
            strand_losses.append(
                {
                    "odf": odf,
                    "cable": cable or "",
                    "port": port,
                    "loss_db": loss,
                    "wavelength": wl,
                    "note": note,
                    "cid": cid,
                }
            )
        elif cable and port:
            strand_losses.append(
                {
                    "odf": "",
                    "cable": cable,
                    "port": port,
                    "loss_db": loss,
                    "wavelength": wl,
                    "note": note,
                    "cid": cid,
                }
            )
        elif cid:
            circuit_losses[cid] = {"loss_db": loss, "wavelength": wl, "note": note}
    return strand_losses, circuit_losses, warnings


def parse_business_sheet(wb: openpyxl.Workbook) -> list[dict[str, str]]:
    ws = wb[SHEET_BUSINESS]
    idx = _header_index(
        ws,
        ["光路编号", "所在机房ODF", "对端ODF", "A端纤芯(排-芯)", "B端纤芯(排-芯)"],
    )
    segments: list[dict[str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        cid = _cell(row, idx, "光路编号")
        dev_a = _cell(row, idx, "所在机房ODF") or _cell(row, idx, "所在机房 ODF")
        dev_b = _cell(row, idx, "对端ODF")
        port_a = _cell(row, idx, "A端纤芯(排-芯)")
        port_b = _cell(row, idx, "B端纤芯(排-芯)")
        if not cid or not dev_a or not dev_b or not port_a or not port_b:
            continue
        segments.append(
            {
                "cid": cid,
                "site": _cell(row, idx, "站点名称"),
                "location": _cell(row, idx, "机房/位置"),
                "cable": _cell(row, idx, "缆段编号"),
                "name": _cell(row, idx, "光路名称"),
                "dev_a": dev_a,
                "port_a": port_a,
                "dev_b": dev_b,
                "port_b": port_b,
                "service": _cell(row, idx, "业务说明"),
                "ip": _cell(row, idx, "IP地址"),
                "label_pos": _cell(row, idx, "标签位置"),
                "note": _cell(row, idx, "备注"),
            }
        )
    return segments


def apply_circuit_losses(segments: list[dict[str, str]], circuit_losses: dict[str, dict[str, str]]) -> None:
    """Attach end-to-end loss from 光损录入 to the last segment of each circuit."""
    if not circuit_losses:
        return
    by_cid: dict[str, list[dict[str, str]]] = {}
    for seg in segments:
        by_cid.setdefault(seg["cid"], []).append(seg)
    for cid, loss_meta in circuit_losses.items():
        segs = by_cid.get(cid)
        if not segs:
            continue
        segs[-1]["loss_db"] = loss_meta["loss_db"]
        segs[-1]["wavelength"] = loss_meta.get("wavelength", "1310")


def validate_business_chain(segments: list[dict[str, str]]) -> list[str]:
    """Check multi-segment circuits connect at ODF+fiber."""
    warnings: list[str] = []
    by_cid: dict[str, list[dict[str, str]]] = {}
    for seg in segments:
        by_cid.setdefault(seg["cid"], []).append(seg)

    for cid, segs in by_cid.items():
        for i in range(len(segs) - 1):
            cur, nxt = segs[i], segs[i + 1]
            if cur["dev_b"] != nxt["dev_a"]:
                warnings.append(
                    f"光路 {cid} 段{i + 1}→{i + 2}: 对端ODF「{cur['dev_b']}」≠ 下段所在ODF「{nxt['dev_a']}」"
                )
            elif cur["port_b"] != nxt["port_a"]:
                warnings.append(
                    f"光路 {cid} 段{i + 1}→{i + 2}: 在 {cur['dev_b']} 纤芯 {cur['port_b']}→{nxt['port_a']} 不衔接"
                )
        for i, seg in enumerate(segs, 1):
            if seg.get("cable") and seg["port_a"] != seg["port_b"]:
                warnings.append(
                    f"光路 {cid} 段{i}: 同缆段 {seg['cable']} 上 A/B 纤芯 {seg['port_a']}≠{seg['port_b']}（通常应相同）"
                )
    return warnings


def parse_workbook(wb: openpyxl.Workbook) -> dict[str, Any]:
    fmt = workbook_format(wb)
    if fmt == "legacy_type":
        return _parse_legacy_type_sheet(wb)
    if fmt != "business_v2":
        raise ValueError("未识别的 Excel 格式，需要「现场登记」业务表或旧版记录类型表")

    base = parse_base_sheet(wb)
    segments = parse_business_sheet(wb)
    strand_losses, circuit_losses, loss_warnings = parse_loss_sheet(wb, segments)
    apply_circuit_losses(segments, circuit_losses)

    return {
        "format": "business_v2",
        "sites": base["sites"],
        "locations": base.get("locations") or [],
        "devices": base["devices"],
        "trunks": base["trunks"],
        "segments": segments,
        "port_links": [],
        "strand_losses": strand_losses,
        "circuit_losses": circuit_losses,
        "warnings": validate_business_chain(segments) + loss_warnings,
    }


def parse_unified_workbook(wb: openpyxl.Workbook) -> dict[str, list[dict[str, str]]]:
    """Compatibility wrapper for import_from_excel."""
    data = parse_workbook(wb)
    return {
        "sites": data["sites"],
        "devices": data["devices"],
        "trunks": data["trunks"],
        "segments": data["segments"],
        "port_links": data["port_links"],
        "strand_losses": data["strand_losses"],
    }


def _parse_legacy_type_sheet(wb: openpyxl.Workbook) -> dict[str, Any]:
    """Old 记录类型* single-sheet format (backward compatible)."""
    ws = wb[SHEET_BUSINESS] if SHEET_BUSINESS in wb.sheetnames else wb.active
    idx = _header_index(ws, ["记录类型"])
    out: dict[str, Any] = {
        "format": "legacy_type",
        "sites": [],
        "devices": [],
        "trunks": [],
        "segments": [],
        "port_links": [],
        "strand_losses": [],
        "circuit_losses": {},
        "warnings": [],
    }
    type_map = {
        "站点": "sites",
        "ODF": "devices",
        "缆段": "trunks",
        "光路": "segments",
        "跳接": "port_links",
    }
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        rtype = _cell(row, idx, "记录类型")
        if not rtype:
            continue
        note = _cell(row, idx, "备注")
        if rtype == "站点":
            name = _cell(row, idx, "站点名称")
            if name:
                out["sites"].append({"name": name, "note": note})
        elif rtype == "ODF":
            name = _cell(row, idx, "ODF设备名")
            if name:
                out["devices"].append(
                    {
                        "name": name,
                        "site": _cell(row, idx, "站点名称"),
                        "location": _cell(row, idx, "机房/位置"),
                        "note": note,
                    }
                )
        elif rtype == "缆段":
            label = _cell(row, idx, "缆段编号")
            da, db = _cell(row, idx, "A端ODF"), _cell(row, idx, "B端ODF")
            if label and da and db:
                out["trunks"].append({"label": label, "dev_a": da, "dev_b": db, "note": note})
        elif rtype == "光路":
            cid = _cell(row, idx, "光路编号")
            da, pa = _cell(row, idx, "A端ODF"), _cell(row, idx, "A端纤芯(排-芯)")
            db, pb = _cell(row, idx, "B端ODF"), _cell(row, idx, "B端纤芯(排-芯)")
            if cid and da and pa and db and pb:
                out["segments"].append(
                    {
                        "cid": cid,
                        "cable": _cell(row, idx, "缆段编号"),
                        "name": _cell(row, idx, "光路名称"),
                        "dev_a": da,
                        "port_a": pa,
                        "dev_b": db,
                        "port_b": pb,
                        "service": _cell(row, idx, "业务说明"),
                        "ip": _cell(row, idx, "IP地址"),
                        "loss_db": _cell(row, idx, "链路光损(dB)"),
                        "wavelength": _cell(row, idx, "链路波长(nm)"),
                        "label_pos": _cell(row, idx, "标签位置"),
                    }
                )
    return out


class ApiClient(Protocol):
    def request(self, method: str, path: str, **kwargs: Any) -> dict[str, Any]: ...
