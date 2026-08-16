#!/usr/bin/env python3
"""Import rooms, ODF, cables from 公司光缆资产_导入确认表.xlsx into NetBox."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import openpyxl

from bootstrap_from_sheet import apply_infra_from_sheet
from fiber_infra import FiberInfra
from import_from_excel import NetBoxClient

DEFAULT_XLSX = Path(__file__).with_name("公司光缆资产_导入确认表.xlsx")
DEFAULT_CONFIG = Path(__file__).with_name("netbox_config.json")


def load_config(path: Path) -> dict:
    with path.open(encoding="utf-8") as f:
        return json.load(f)


def _parse_trunk_note(note: str) -> tuple[str, str]:
    note = (note or "").strip()
    strands, length = "", ""
    m = re.search(r"芯数\s*(\d+)", note)
    if m:
        strands = m.group(1)
    m = re.search(r"长度\s*(\d+)\s*m", note)
    if m:
        length = m.group(1)
    return strands, length


def parse_workbook(xlsx: Path) -> dict[str, list[dict[str, str]]]:
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    if "基础配置" not in wb.sheetnames:
        raise SystemExit(f"{xlsx} 缺少「基础配置」工作表，请先运行 parse_asset_sheet.py")

    segment_meta: dict[str, dict[str, str]] = {}
    if "光缆缆段" in wb.sheetnames:
        ws_seg = wb["光缆缆段"]
        for row in ws_seg.iter_rows(min_row=2, values_only=True):
            if not row or not row[3]:
                continue
            label = str(row[3]).strip()
            cores = str(row[4]).strip() if row[4] is not None else ""
            length = str(row[5]).strip() if len(row) > 5 and row[5] is not None else ""
            if cores.endswith("芯"):
                cores = cores.replace("芯", "")
            segment_meta[label] = {"strands": cores, "length": length}

    sites: list[dict[str, str]] = []
    locations: list[dict[str, str]] = []
    devices: list[dict[str, str]] = []
    trunks: list[dict[str, str]] = []

    ws = wb["基础配置"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        kind = str(row[0]).strip()
        name = str(row[1] or "").strip()
        rel_a = str(row[2] or "").strip()
        rel_b = str(row[3] or "").strip()
        note = str(row[4] or "").strip() if len(row) > 4 else ""
        if not name:
            continue
        if kind == "站点":
            sites.append({"name": name, "slug": rel_a, "note": note})
        elif kind == "机房":
            locations.append({"name": name, "site": rel_a, "note": note})
        elif kind == "ODF":
            devices.append({"name": name, "site": rel_a, "location": rel_b, "note": note})
        elif kind == "缆段":
            meta = segment_meta.get(name, {})
            strands = meta.get("strands", "")
            length = meta.get("length", "")
            if not strands or not length:
                ns, nl = _parse_trunk_note(note)
                strands = strands or ns
                length = length or nl
            trunks.append(
                {
                    "label": name,
                    "dev_a": rel_a,
                    "dev_b": rel_b,
                    "note": note,
                    "strands": strands or "12",
                    "length": length,
                }
            )
    wb.close()
    return {"sites": sites, "locations": locations, "devices": devices, "trunks": trunks}


def main() -> None:
    ap = argparse.ArgumentParser(description="从资产确认表导入 NetBox 机房/ODF/缆段")
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    ap.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--skip-ports", action="store_true", help="仅创建设备与缆段，不 provision 端口")
    args = ap.parse_args()

    if not args.xlsx.exists():
        raise SystemExit(f"找不到 {args.xlsx}，请先运行: python parse_asset_sheet.py")
    if not args.config.exists():
        raise SystemExit(f"找不到 {args.config}")

    data = parse_workbook(args.xlsx)
    cfg = load_config(args.config)
    client = NetBoxClient(cfg["base_url"], cfg["token"], cfg.get("verify_ssl", True))

    apply_infra_from_sheet(
        client,
        data["sites"],
        data["devices"],
        data["trunks"],
        locations=data["locations"],
        dry_run=args.dry_run,
    )

    if args.dry_run or args.skip_ports:
        return

    infra = FiberInfra(client)
    print("\n==> 配置 ODF 端口（站点缩写 to 站点缩写 + 排-芯）")
    for row in data["trunks"]:
        label = row["label"]
        for dev in (row["dev_a"], row["dev_b"]):
            try:
                infra.provision_cable(dev, label)
                infra.fix_device_port_names(dev)
            except Exception as exc:
                print(f"  ⚠ {dev} @ {label}: {exc}")
        try:
            infra.fix_cable_termination(label)
        except Exception as exc:
            print(f"  ⚠ 改接 {label}: {exc}")
    print("完成。")


if __name__ == "__main__":
    main()
