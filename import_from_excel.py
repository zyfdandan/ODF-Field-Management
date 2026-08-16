#!/usr/bin/env python3
"""Import multi-segment fiber circuits from Excel; one 光路编号 = one logical link."""

from __future__ import annotations

import argparse
import json
import sys
from collections import defaultdict
from pathlib import Path
from typing import TYPE_CHECKING, Any

import requests

if TYPE_CHECKING:
    import openpyxl

from bootstrap_from_sheet import apply_infra_from_sheet
from unified_sheet import SHEET_BUSINESS, is_unified_workbook, parse_workbook
from fiber_infra import (
    FiberInfra,
    infer_links_from_segments,
    infer_splices_from_segments,
    merge_link_rows,
    merge_splice_rows,
    resolve_port_name,
    splice_rows_to_links,
)

DEFAULT_CONFIG = Path(__file__).with_name("netbox_config.json")
DEFAULT_EXCEL = Path(__file__).with_name("现场登记.xlsx")
DEFAULT_EXCEL_LEGACY = Path(__file__).with_name("现场登记表.xlsx")

COL = {
    "cid": "光路编号",
    "seg": "段序号",
    "name": "光路名称",
    "cable": "缆段编号",
    "dev_a": "A端设备",
    "port_a": "A端纤芯(排-芯)",
    "dev_b": "B端设备",
    "port_b": "B端纤芯(排-芯)",
    "service": "业务说明",
    "ip": "IP地址",
    "loss_db": "实测光损(dB)",
    "wavelength": "测试波长(nm)",
    "label_pos": "标签位置",
    "url": "扫码链接",
    "qr_file": "二维码",
}

SPLICE_COL = {
    "dev": "熔接点设备",
    "plan": "熔接方案名称",
    "cab_a": "A端缆段编号",
    "port_a": "A端纤芯(排-芯)",
    "cab_b": "B端缆段编号",
    "port_b": "B端纤芯(排-芯)",
    "tray": "托盘序号",
    "note": "备注",
    "cid": "关联光路编号",
}

PORT_LINK_COL = {
    "dev_a": "A端设备",
    "cab_a": "A端缆段编号",
    "port_a": "A端纤芯(排-芯)",
    "dev_b": "B端设备",
    "cab_b": "B端缆段编号",
    "port_b": "B端纤芯(排-芯)",
    "note": "备注",
    "cid": "关联光路编号",
}

STRAND_LOSS_COL = {
    "cable": "缆段编号",
    "port": "纤芯(排-芯)",
    "loss_db": "实测光损(dB)",
    "wavelength": "测试波长(nm)",
    "note": "备注",
    "cid": "关联光路编号",
}


def normalize_cell(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, float):
        if val == int(val):
            return str(int(val))
        return format(val, "g")
    return str(val).strip()


class NetBoxClient:
    def __init__(self, base_url: str, token: str, verify_ssl: bool = True) -> None:
        self.base_url = base_url.rstrip("/")
        self.session = requests.Session()
        self.session.headers.update(
            {
                "Authorization": f"Token {token}",
                "Accept": "application/json",
                "Content-Type": "application/json",
            }
        )
        self.session.verify = verify_ssl

    def request(self, method: str, path: str, **kwargs: Any) -> dict[str, Any]:
        url = f"{self.base_url}{path}"
        resp = self.session.request(method, url, **kwargs)
        if resp.status_code >= 400:
            raise RuntimeError(f"{method} {path} -> {resp.status_code}: {resp.text[:800]}")
        if resp.status_code == 204 or not resp.content:
            return {}
        return resp.json()

    def post(self, path: str, data: dict[str, Any]) -> dict[str, Any]:
        return self.request("POST", path, json=data)


def read_segments(wb: openpyxl.Workbook) -> list[dict[str, str]]:
    ws = wb["光路录入"] if "光路录入" in wb.sheetnames else wb.active
    headers = [str(c.value or "").replace("*", "") for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    for r in (COL["cid"], COL["dev_a"], COL["port_a"], COL["dev_b"], COL["port_b"]):
        if r not in idx:
            raise ValueError(f"缺少列: {r}")

    rows: list[dict[str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[idx[COL["cid"]]]:
            continue
        item = {
            key: normalize_cell(row[idx[COL[key]]])
            for key in COL
            if key
            in (
                "cid",
                "seg",
                "name",
                "dev_a",
                "port_a",
                "dev_b",
                "port_b",
                "service",
                "ip",
                "loss_db",
                "wavelength",
                "label_pos",
            )
            and COL[key] in idx
        }
        if COL["cable"] in idx:
            item["cable"] = normalize_cell(row[idx[COL["cable"]]])
        if not all(item.get(k) for k in ("cid", "dev_a", "port_a", "dev_b", "port_b")):
            continue
        rows.append(item)
    return rows


def read_strand_losses(wb: openpyxl.Workbook) -> list[dict[str, str]]:
    if "缆段纤芯光损" not in wb.sheetnames:
        return []
    ws = wb["缆段纤芯光损"]
    headers = [str(c.value or "").replace("*", "") for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    for r in (STRAND_LOSS_COL["cable"], STRAND_LOSS_COL["port"], STRAND_LOSS_COL["loss_db"]):
        if r not in idx:
            raise ValueError(f"缆段纤芯光损缺少列: {r}")

    rows: list[dict[str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[idx[STRAND_LOSS_COL["cable"]]]:
            continue
        item = {
            key: normalize_cell(row[idx[STRAND_LOSS_COL[key]]])
            for key in STRAND_LOSS_COL
            if STRAND_LOSS_COL[key] in idx
        }
        if not all(item.get(k) for k in ("cable", "port", "loss_db")):
            continue
        rows.append(item)
    return rows


def read_port_links(wb: openpyxl.Workbook) -> list[dict[str, str]]:
    if "端口连接" not in wb.sheetnames:
        return []
    ws = wb["端口连接"]
    headers = [str(c.value or "").replace("*", "") for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    for r in (PORT_LINK_COL["dev_a"], PORT_LINK_COL["port_a"], PORT_LINK_COL["dev_b"], PORT_LINK_COL["port_b"]):
        if r not in idx:
            raise ValueError(f"端口连接缺少列: {r}")

    rows: list[dict[str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[idx[PORT_LINK_COL["dev_a"]]]:
            continue
        item = {
            key: normalize_cell(row[idx[PORT_LINK_COL[key]]])
            for key in PORT_LINK_COL
            if PORT_LINK_COL[key] in idx
        }
        if not all(item.get(k) for k in ("dev_a", "port_a", "dev_b", "port_b")):
            continue
        rows.append(item)
    return rows


def read_splices(wb: openpyxl.Workbook) -> list[dict[str, str]]:
    if "熔接录入" not in wb.sheetnames:
        return []
    ws = wb["熔接录入"]
    headers = [str(c.value or "").replace("*", "") for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    for r in (SPLICE_COL["dev"], SPLICE_COL["port_a"], SPLICE_COL["port_b"]):
        if r not in idx:
            raise ValueError(f"熔接录入缺少列: {r}")

    rows: list[dict[str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[idx[SPLICE_COL["dev"]]]:
            continue
        item = {
            key: normalize_cell(row[idx[SPLICE_COL[key]]])
            for key in SPLICE_COL
            if SPLICE_COL[key] in idx
        }
        if not all(item.get(k) for k in ("dev", "port_a", "port_b")):
            continue
        if not item.get("cab_a") or not item.get("cab_b"):
            raise ValueError(
                f"熔接行 {item['dev']} {item['port_a']}↔{item['port_b']} 缺少 A/B 端缆段编号"
            )
        rows.append(item)
    return rows


def auto_order_segments(segments: list[dict[str, str]]) -> list[dict[str, str]]:
    """Order segments: by 段序号 if all filled, else topological chain by A/B match."""
    if all(s.get("seg") for s in segments):
        return sorted(segments, key=lambda s: int(s["seg"]))

    if len(segments) == 1:
        return segments

    # endpoints: dev_a that is never someone's dev_b = chain start
    all_a = {s["dev_a"] for s in segments}
    all_b = {s["dev_b"] for s in segments}
    starts = [s for s in segments if s["dev_a"] not in all_b]
    ends = [s for s in segments if s["dev_b"] not in all_a]

    if len(starts) == 1:
        ordered: list[dict[str, str]] = []
        current = starts[0]
        remaining = segments.copy()
        while current:
            ordered.append(current)
            remaining.remove(current)
            nxt = [
                s
                for s in remaining
                if s["dev_a"] == current["dev_b"] and s["port_a"] == current["port_b"]
            ]
            if not nxt:
                nxt = [s for s in remaining if s["dev_a"] == current["dev_b"]]
            if len(nxt) > 1:
                nxt = sorted(nxt, key=lambda s: (s["port_a"] != current["port_b"], s.get("cable") or ""))
            current = nxt[0] if nxt else None
        if len(ordered) == len(segments):
            return ordered

    if len(starts) == 1 and len(ends) == 1:
        # relax port match fallback already tried; warn and use sheet order
        pass

    print(f"  提示: 无法唯一排序，按表格顺序组合（共 {len(segments)} 段）")
    return segments


def group_by_circuit(rows: list[dict[str, str]]) -> dict[str, list[dict[str, str]]]:
    groups: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        groups[row["cid"]].append(row)
    for cid in groups:
        groups[cid] = auto_order_segments(groups[cid])
        for i, seg in enumerate(groups[cid], start=1):
            seg["_order"] = str(i)
    return dict(groups)


def validate_chain(segments: list[dict[str, str]]) -> list[str]:
    warnings: list[str] = []
    for s in segments:
        if s["port_a"] != s["port_b"]:
            warnings.append(
                f"  提示: 段{s.get('_order', s.get('seg', '?'))} "
                f"A纤芯({s['port_a']}) ≠ B纤芯({s['port_b']})，"
                f"同一段缆线通常应一致（熔接/改道时除外）"
            )
    for i in range(len(segments) - 1):
        s1, s2 = segments[i], segments[i + 1]
        if s1["dev_b"] != s2["dev_a"]:
            warnings.append(
                f"  警告: 段{s1.get('_order', s1.get('seg', '?'))} B端({s1['dev_b']}) "
                f"≠ 段{s2.get('_order', s2.get('seg', '?'))} A端({s2['dev_a']})"
            )
        if s1["port_b"] != s2["port_a"]:
            warnings.append(
                f"  警告: 段{s1.get('_order', s1.get('seg', '?'))} B纤芯({s1['port_b']}) "
                f"≠ 段{s2.get('_order', s2.get('seg', '?'))} A纤芯({s2['port_a']})"
            )
    return warnings


def merge_circuit_meta(segments: list[dict[str, str]]) -> dict[str, str]:
    first = segments[0]
    name = first.get("name") or next((s.get("name") for s in segments if s.get("name")), first["cid"])
    return {
        "name": name,
        "ip": next((s.get("ip") for s in segments if s.get("ip")), ""),
        "service": next((s.get("service") for s in segments if s.get("service")), ""),
        "loss_db": next((s.get("loss_db") for s in reversed(segments) if s.get("loss_db")), ""),
        "wavelength": next((s.get("wavelength") for s in reversed(segments) if s.get("wavelength")), ""),
        "label_pos": " | ".join(s.get("label_pos") for s in segments if s.get("label_pos")),
    }


def build_table_description(segments: list[dict[str, str]], meta: dict[str, str]) -> str:
    lines = [f"业务IP: {meta.get('ip') or '-'}", f"逻辑段数: {len(segments)}", "—— 表格段明细 ——"]
    for s in segments:
        seg_no = s.get("_order") or s.get("seg") or "?"
        cab = s.get("cable") or "-"
        ip = f" IP:{s['ip']}" if s.get("ip") else ""
        loss = ""
        if s.get("loss_db"):
            wl = s.get("wavelength") or "1310"
            loss = f" 光损:{s['loss_db']}dB@{wl}nm"
        lines.append(
            f"段{seg_no} [{cab}] {s['dev_a']}:{s['port_a']} → {s['dev_b']}:{s['port_b']}"
            f" | {s.get('service') or '-'}{ip}{loss}"
        )
    return "\n".join(lines)


def build_segment_summary(segments: list[dict[str, str]]) -> str:
    parts = []
    for i, s in enumerate(segments, start=1):
        seg_no = s.get("_order") or s.get("seg") or str(i)
        parts.append(f"段{seg_no}:{s['dev_a']}:{s['port_a']}→{s['dev_b']}:{s['port_b']}")
    return " → ".join(parts)


def build_route_short(segments: list[dict[str, str]]) -> str:
    """Short human route: ODF1 → ODF2 → ODF3"""
    nodes = [segments[0]["dev_a"]]
    for s in segments:
        nodes.append(s["dev_b"])
    # dedupe consecutive same
    short = [nodes[0]]
    for n in nodes[1:]:
        if n != short[-1]:
            short.append(n)
    return " → ".join(short)


def build_odf_label_manifest(segments: list[dict[str, str]], cid: str, name: str) -> list[dict[str, str]]:
    route_short = build_route_short(segments)
    n_seg = len(segments)
    labels: list[dict[str, str]] = []
    seen: set[tuple[str, str, str, str]] = set()

    for s in segments:
        seg_no = s.get("_order") or s.get("seg") or "?"
        for dev, peer, fiber in (
            (s["dev_a"], s["dev_b"], s["port_a"]),
            (s["dev_b"], s["dev_a"], s["port_b"]),
        ):
            key = (dev, peer, fiber, seg_no)
            if key in seen:
                continue
            seen.add(key)
            labels.append(
                {
                    "光路编号": cid,
                    "光路名称": name,
                    "段序号": seg_no,
                    "全程段数": str(n_seg),
                    "全程路由": route_short,
                    "本端ODF": dev,
                    "本端纤芯": fiber,
                    "对端ODF(本段)": peer,
                    "ODF框标签建议": (
                        f"{dev}\nODF\n对端：{peer}\n"
                        f"全程：{route_short}（{n_seg}段）\n{cid}"
                    ),
                }
            )
    return labels


def find_device(client: NetBoxClient, name: str) -> dict[str, Any]:
    name = str(name or "").strip()
    if not name or name.lower() in ("undefined", "null", "none"):
        raise RuntimeError("找不到设备: 名称为空，请确认对端 ODF 已解析")
    hits = client.request("GET", "/dcim/devices/", params={"name": name, "limit": 10})
    results = hits.get("results", [])
    if not results:
        raise RuntimeError(f"找不到设备: {name}")
    if len(results) > 1:
        raise RuntimeError(f"设备名重复: {name}")
    return results[0]


def find_front_port(client: NetBoxClient, device_id: int, port_name: str, fallback: str | None = None) -> dict[str, Any]:
    hits = client.request(
        "GET",
        "/dcim/front-ports/",
        params={"device_id": device_id, "name": port_name, "limit": 10},
    )
    results = hits.get("results", [])
    if not results and fallback and fallback != port_name:
        hits = client.request(
            "GET",
            "/dcim/front-ports/",
            params={"device_id": device_id, "name": fallback, "limit": 10},
        )
        results = hits.get("results", [])
    if not results:
        raise RuntimeError(f"设备 id={device_id} 上找不到纤芯: {port_name}")
    return results[0]


def _path_optical_fields(meta: dict[str, str]) -> dict[str, Any]:
    payload: dict[str, Any] = {}
    loss = (meta.get("loss_db") or "").strip()
    if loss:
        payload["actual_loss_db"] = loss
        wl = (meta.get("wavelength") or "").strip()
        payload["wavelength_nm"] = int(wl) if wl else 1310
    elif (meta.get("wavelength") or "").strip():
        payload["wavelength_nm"] = int(meta["wavelength"])
    return payload


def _write_back_excel(
    wb: openpyxl.Workbook,
    excel_path: Path,
    results: list[dict[str, str]],
    col_url: int,
    col_qr: int,
) -> None:
    from excel_qr import write_qr_back_to_excel

    if SHEET_BUSINESS in wb.sheetnames:
        ws = wb[SHEET_BUSINESS]
    elif "光路录入" in wb.sheetnames:
        ws = wb["光路录入"]
    else:
        ws = wb.active
    headers = [str(c.value or "").replace("*", "") for c in ws[1]]
    cid_col = headers.index(COL["cid"]) + 1
    n = write_qr_back_to_excel(
        wb,
        excel_path,
        cid_col=cid_col,
        col_url=col_url,
        col_qr=col_qr,
        results=results,
        sheet_name=ws.title,
    )
    if n:
        print(f"  已嵌入 {n} 个二维码到 Excel「二维码」列")


def import_circuits(
    excel_path: Path,
    config: dict[str, Any],
    dry_run: bool = False,
    skip_infra: bool = False,
    link_mode: str = "port",
    sync_osp: bool = True,
) -> list[dict[str, str]]:
    import openpyxl

    wb = openpyxl.load_workbook(excel_path)
    unified = is_unified_workbook(wb)

    if is_unified_workbook(wb):
        parsed = parse_workbook(wb)
        rows = parsed["segments"]
        manual_port_links = parsed.get("port_links") or []
        manual_splices = [] if link_mode == "port" else []
        strand_losses = parsed.get("strand_losses") or []
        fmt = parsed.get("format", "business_v2")
        print(
            f"{'三表' if fmt == 'business_v2' else '单表'}模式: "
            f"基础(站点 {len(parsed['sites'])} / ODF {len(parsed['devices'])} / 缆段 {len(parsed['trunks'])}) | "
            f"业务 {len(rows)} 段 | 光损 {len(strand_losses)} 条 | 跳接 {len(manual_port_links)} 条"
        )
        for w in parsed.get("warnings") or []:
            print(f"  ⚠ {w}")
    else:
        rows = read_segments(wb)
        manual_port_links = read_port_links(wb)
        manual_splices = read_splices(wb) if link_mode == "splice" else []
        strand_losses = read_strand_losses(wb)
        parsed = None

    if link_mode == "port" and not manual_port_links and not unified:
        manual_port_links = splice_rows_to_links(read_splices(wb))
    if not rows:
        print("没有可导入的数据行。")
        return []

    groups = group_by_circuit(rows)
    splices: list[dict[str, str]] = []
    links: list[dict[str, str]] = []
    if link_mode == "port":
        inferred_links = infer_links_from_segments(groups)
        links = merge_link_rows(manual_port_links, inferred_links)
        if inferred_links:
            print(
                f"端口连接: 表格 {len(manual_port_links)} 条 + 自动段间 {len(inferred_links)} 条 = 共 {len(links)} 条"
            )
        elif manual_port_links:
            print(f"端口连接: 表格 {len(manual_port_links)} 条")
    else:
        inferred_splices = infer_splices_from_segments(groups)
        splices = merge_splice_rows(manual_splices, inferred_splices)
        if inferred_splices:
            print(f"熔接: 表格 {len(manual_splices)} 条 + 自动段间 {len(inferred_splices)} 条 = 共 {len(splices)} 条")

    cable_labels: list[str] = []
    for seg in rows:
        if seg.get("cable"):
            cable_labels.append(seg["cable"])
    for sp in splices:
        if sp.get("cab_a"):
            cable_labels.append(sp["cab_a"])
        if sp.get("cab_b"):
            cable_labels.append(sp["cab_b"])
    for lk in links:
        if lk.get("cab_a"):
            cable_labels.append(lk["cab_a"])
        if lk.get("cab_b"):
            cable_labels.append(lk["cab_b"])

    client = NetBoxClient(config["base_url"], config["token"], config.get("verify_ssl", True))
    web_base = config.get("web_base_url", config["base_url"].replace("/api", ""))
    qr_dir = Path(config.get("qr_output_dir", excel_path.parent / "qr_labels"))
    qr_dir.mkdir(parents=True, exist_ok=True)

    if unified and parsed and not skip_infra:
        apply_infra_from_sheet(
            client,
            parsed["sites"],
            parsed["devices"],
            parsed["trunks"],
            locations=parsed.get("locations") or [],
            dry_run=dry_run,
        )

    infra = None if dry_run else FiberInfra(client)
    if infra and not skip_infra:
        infra.apply_from_excel(rows, splices=splices, links=links, link_mode=link_mode)
        if strand_losses:
            print(f"\n==> ODF框光损录入 ({len(strand_losses)} 行)")
            n = infra.apply_strand_losses(strand_losses)
            print(f"  已写入 {n} 条纤芯实测光损")
    elif link_mode == "port" and not links and not dry_run:
        print("提示: 未填写「端口连接」，多缆 ODF 站点可能无法全链路连通。")
    elif link_mode == "splice" and not splices and not dry_run:
        print("提示: 未填写「熔接录入」，多缆 ODF 站点可能无法全链路连通。")

    ws = wb[SHEET_BUSINESS] if unified and SHEET_BUSINESS in wb.sheetnames else (
        wb["光路录入"] if "光路录入" in wb.sheetnames else wb.active
    )
    headers = [str(c.value or "").replace("*", "") for c in ws[1]]
    col_url = headers.index("扫码链接") + 1 if "扫码链接" in headers else (
        headers.index(COL["url"]) + 1 if COL["url"] in headers else None
    )
    from excel_qr import find_qr_column, save_qr_png

    col_qr = find_qr_column(headers)

    results: list[dict[str, str]] = []
    all_odf_labels: list[dict[str, str]] = []
    expected_endpoints: dict[str, tuple[int, int, str, str]] = {}

    for cid, segments in groups.items():
        print(f"\n==> {cid}（{len(segments)} 段，自动组合）")
        print(f"  全程: {build_route_short(segments)}")
        for w in validate_chain(segments):
            print(w)

        meta = merge_circuit_meta(segments)
        first, last = segments[0], segments[-1]
        table_desc = build_table_description(segments, meta)
        port_a_name = resolve_port_name(first["port_a"], first.get("cable", ""), device_name=first["dev_a"])
        port_b_name = resolve_port_name(last["port_b"], last.get("cable", ""), device_name=last["dev_b"])

        desc_parts = [table_desc]
        if meta.get("loss_db"):
            wl = meta.get("wavelength") or "1310"
            desc_parts.append(f"全链路光损: {meta['loss_db']}dB@{wl}nm")
        comments = table_desc
        if meta.get("loss_db"):
            wl = meta.get("wavelength") or "1310"
            comments += f"\n全链路光损: {meta['loss_db']}dB@{wl}nm"
        circuit_payload = {
            "name": meta["name"],
            "cid": cid,
            "status": "active",
            "strand_count": 1,
            "description": "\n".join(desc_parts)[:2000],
            "comments": comments[:4000],
        }

        print(f"  表格段数: {len(segments)} | 业务IP: {meta.get('ip') or '-'}")
        print(
            f"  端到端: {first['dev_a']}:{first['port_a']}({port_a_name}) → "
            f"… → {last['dev_b']}:{last['port_b']}({port_b_name})"
        )
        for i, s in enumerate(segments, 1):
            print(f"    段{i} [{s.get('cable') or '-'}] {s['dev_a']}:{s['port_a']} → {s['dev_b']}:{s['port_b']}")

        if dry_run:
            print(f"  DRY-RUN: {circuit_payload}")
            all_odf_labels.extend(build_odf_label_manifest(segments, cid, meta.get("name", "")))
            continue

        dev_a = find_device(client, first["dev_a"])
        dev_b = find_device(client, last["dev_b"])
        port_a = find_front_port(client, dev_a["id"], port_a_name, first["port_a"])
        port_b = find_front_port(client, dev_b["id"], port_b_name, last["port_b"])
        expected_endpoints[cid] = (port_a["id"], port_b["id"], last["dev_b"], last["port_b"])

        existing = client.request("GET", "/plugins/fms/fiber-circuits/", params={"cid": cid, "limit": 1})
        if existing.get("results"):
            circuit_id = existing["results"][0]["id"]
            client.request("PATCH", f"/plugins/fms/fiber-circuits/{circuit_id}/", json=circuit_payload)
        else:
            circuit = client.post("/plugins/fms/fiber-circuits/", circuit_payload)
            circuit_id = circuit["id"]

        paths = client.request(
            "GET", "/plugins/fms/fiber-circuit-paths/", params={"circuit_id": circuit_id, "limit": 10}
        )
        path_payload = {
            "circuit": circuit_id,
            "position": 1,
            "origin": port_a["id"],
            "destination": port_b["id"],
            "path": [{"type": "front_port", "id": port_a["id"]}],
            **_path_optical_fields(meta),
        }
        if paths.get("results"):
            path_id = paths["results"][0]["id"]
            client.request("PATCH", f"/plugins/fms/fiber-circuit-paths/{path_id}/", json=path_payload)
        else:
            path = client.post("/plugins/fms/fiber-circuit-paths/", path_payload)
            path_id = path["id"]

        client.request("POST", f"/plugins/fms/fiber-circuits/{circuit_id}/retrace/")
        client.request(
            "PATCH",
            f"/plugins/fms/fiber-circuit-paths/{path_id}/",
            json={"origin": port_a["id"], "destination": port_b["id"]},
        )
        paths_after = client.request(
            "GET", "/plugins/fms/fiber-circuit-paths/", params={"circuit_id": circuit_id, "limit": 5}
        )
        for p in paths_after.get("results", []):
            if p["id"] == path_id:
                ok = p.get("is_complete")
                hops = len(p.get("path") or [])
                dest = p.get("destination")
                dest_id = dest["id"] if isinstance(dest, dict) else dest
                dest_fp = client.request("GET", f"/dcim/front-ports/{dest_id}/") if dest_id else {}
                dest_label = (
                    f"{dest_fp.get('device', {}).get('name', '?')}:{dest_fp.get('name', '?')}"
                    if dest_fp
                    else "?"
                )
                print(f"  路径: {'完整' if ok else '未完整'}, hops={hops}, 终点={dest_label}")
                if not ok or dest_id != port_b["id"]:
                    print(f"  ⚠ 期望终点 {last['dev_b']}:{last['port_b']} — 请检查熔接/缆段编号")
                elif infra and ok:
                    chain, chain_detail = infra.apply_fiber_chain(path_id)
                    if chain:
                        print(f"  完整纤芯链: {chain}")
                        chain_block = f"完整纤芯链: {chain}"
                        if chain_detail:
                            chain_block += f"\n—— 逐站纤芯 ——\n{chain_detail}"
                        client.request(
                            "PATCH",
                            f"/plugins/fms/fiber-circuits/{circuit_id}/",
                            json={
                                "description": f"{table_desc}\n{chain_block}"[:2000],
                                "comments": f"{comments}\n{chain_block}"[:4000],
                            },
                        )
        trace_url = f"{web_base}/plugins/fms/fiber-circuit-paths/{path_id}/#trace"
        qr_png = qr_dir / f"{cid}.png"
        qr_file = f"qr_labels/{cid}.png"
        if save_qr_png(qr_png, trace_url):
            results.append(
                {
                    "cid": cid,
                    "url": trace_url,
                    "qr_file": qr_file,
                    "qr_abs_path": str(qr_png.resolve()),
                }
            )
        else:
            results.append({"cid": cid, "url": trace_url, "qr_file": "", "qr_abs_path": ""})
        all_odf_labels.extend(build_odf_label_manifest(segments, cid, meta["name"]))

    if not dry_run and results:
        if col_url and col_qr:
            _write_back_excel(wb, excel_path, results, col_url, col_qr)
        (qr_dir / "qr_manifest.json").write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
        (qr_dir / "odf_labels.json").write_text(
            json.dumps(all_odf_labels, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        print(f"\nODF 标签建议: {qr_dir / 'odf_labels.json'}")

    print(f"\n完成，共 {len(results)} 条逻辑光路（{len(rows)} 段录入行，{len(splices)} 条熔接）。")

    if not dry_run and results and not skip_infra and infra:
        failures = infra.retrace_and_verify([r["cid"] for r in results], expected_endpoints)
        if failures:
            print("\n以下光路未完整连通：")
            for f in failures:
                print(f"  - {f}")
            raise SystemExit(1)
        print("\n全链路校验通过。")

    if sync_osp and not dry_run and not skip_infra and results:
        from sync_osp_strands import sync_osp_from_dcim

        print("\n==> 同步 OSP Strand（FrontPort 追踪）")
        stats = sync_osp_from_dcim(client)
        print(
            f"  OSP: 修正 {stats.get('repaired', 0)} 条, {stats['trunks']} 缆, {stats['strands']} Strand, "
            f"{stats['jumps']} 跳线, +{stats['splices']} 熔接"
        )

    return results


def main() -> None:
    parser = argparse.ArgumentParser(description="Import multi-segment fiber circuits from Excel")
    parser.add_argument("--excel", type=Path, default=DEFAULT_EXCEL)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--skip-infra", action="store_true", help="跳过端口/连接/缆线配置（仅导入光路）")
    parser.add_argument(
        "--link-mode",
        choices=("port", "splice"),
        default="port",
        help="port=FrontPort 对 FrontPort DCIM 连接（默认）；splice=熔接方案+FMS 熔接条目",
    )
    parser.add_argument(
        "--no-sync-osp",
        action="store_true",
        help="跳过 netbox-osp Strand/Splice 同步（FrontPort OSP 追踪将不可用）",
    )
    args = parser.parse_args()

    if not args.config.exists():
        print(f"缺少配置: {args.config}", file=sys.stderr)
        sys.exit(1)
    excel = args.excel
    if not excel.exists():
        for alt in (
            DEFAULT_EXCEL_LEGACY,
            Path(__file__).with_name("现场登记表.xlsx"),
            Path(__file__).with_name("现场光路录入.xlsx"),
            excel.with_name("现场登记_新.xlsx"),
            excel.with_name("现场登记表_新.xlsx"),
            excel.with_name("现场光路录入_新.xlsx"),
        ):
            if alt.exists():
                excel = alt
                break
    if not excel.exists():
        print(f"缺少 Excel: {args.excel}", file=sys.stderr)
        sys.exit(1)

    config = json.loads(args.config.read_text(encoding="utf-8"))
    import_circuits(
        excel,
        config,
        dry_run=args.dry_run,
        skip_infra=args.skip_infra,
        link_mode=args.link_mode,
        sync_osp=not args.no_sync_osp,
    )


if __name__ == "__main__":
    main()
